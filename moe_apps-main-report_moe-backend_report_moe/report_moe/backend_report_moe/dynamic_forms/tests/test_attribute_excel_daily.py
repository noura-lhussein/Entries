"""Attribute-driven Excel for daily national titles (Form Builder path)."""

from __future__ import annotations

import io
from pathlib import Path

from django.contrib.auth import get_user_model
from django.core.management import call_command
from django.test import TestCase
from openpyxl import load_workbook

from dynamic_forms.models import Attribute, Info, MainSection, SubMainSection, Title
from dynamic_forms.title_excel.dynamic import DynamicTitleExcelHandler
from dynamic_forms.title_excel.electricity_daily import ElectricityDailyExcelHandler
from dynamic_forms.title_excel.oil_gas_daily import OilGasDailyExcelHandler
from dynamic_forms.title_excel.measures import COL_KEY, SHEET_METRICS
from dynamic_forms.title_excel.registry import get_handler_for_title
from dynamic_forms.title_excel.schema import importable_attributes
from electricity.report_template import EXPORT_SHEET_ORDER
from oil_gas.report_template import EXPORT_SHEET_ORDER as OIL_GAS_SHEET_ORDER

User = get_user_model()
FIXTURE = (
    Path(__file__).resolve().parent
    / "fixtures"
    / "electricity_daily_report_ref.xlsx"
)
OIL_FIXTURE = (
    Path(__file__).resolve().parent
    / "fixtures"
    / "oil_gas_daily_report_ref.xlsx"
)


class AttributeExcelDailyTests(TestCase):
    @classmethod
    def setUpTestData(cls):
        call_command("seed_oil_gas_daily_info_forms", verbosity=0)
        call_command("seed_electricity_daily_info_forms", verbosity=0)
        try:
            call_command("sync_info_report_date_flag", "--apply", verbosity=0)
        except Exception:
            pass
        Attribute.objects.filter(key="report_date").update(is_report_date=True)

        cls.oil_title = Title.objects.get(code="oil_gas.national")
        cls.el_title = Title.objects.get(code="electricity.national")
        if getattr(cls.oil_title, "entry_mode", None) == "multi_record":
            Title.objects.filter(pk=cls.oil_title.pk).update(
                entry_mode="single_record")
            cls.oil_title.refresh_from_db()
        if getattr(cls.el_title, "entry_mode", None) == "multi_record":
            Title.objects.filter(pk=cls.el_title.pk).update(
                entry_mode="single_record")
            cls.el_title.refresh_from_db()

        main = MainSection.objects.create(name="Test Main Excel")
        cls.sub_main = SubMainSection.objects.create(
            name="Test Sub Excel", main_section=main
        )
        cls.user = User.objects.create_user(
            email="excel-daily@test.local",
            password="Test@2026!",
        )
        cls.user.is_staff = True
        cls.user.save(update_fields=["is_staff"])
        cls.handler = DynamicTitleExcelHandler()
        cls.el_handler = ElectricityDailyExcelHandler()
        cls.og_handler = OilGasDailyExcelHandler()

    def _fill_template(self, title: Title, values_by_key: dict[str, str]) -> io.BytesIO:
        raw = self.handler.build_template_bytes(
            title, sub_main_id=self.sub_main.id)
        wb = load_workbook(io.BytesIO(raw))
        ws = wb[SHEET_METRICS]
        headers = [c.value for c in ws[1]]
        key_idx = headers.index(COL_KEY) + 1
        value_idx = 3
        for row in range(2, ws.max_row + 1):
            key = ws.cell(row=row, column=key_idx).value
            if key and str(key).strip() in values_by_key:
                ws.cell(
                    row=row,
                    column=value_idx,
                    value=values_by_key[str(key).strip()],
                )
        buf = io.BytesIO()
        wb.save(buf)
        buf.seek(0)
        return buf

    def test_registry_routes_sector_nationals(self):
        self.assertIsInstance(
            get_handler_for_title(self.el_title), ElectricityDailyExcelHandler
        )
        self.assertIsInstance(
            get_handler_for_title(self.oil_title), OilGasDailyExcelHandler
        )

    def test_electricity_template_matches_official_sheet_order(self):
        raw = self.el_handler.build_template_bytes(self.el_title)
        wb = load_workbook(io.BytesIO(raw))
        self.assertEqual(tuple(wb.sheetnames), EXPORT_SHEET_ORDER)

    def test_oil_gas_template_matches_official_sheet_order(self):
        raw = self.og_handler.build_template_bytes(self.oil_title)
        wb = load_workbook(io.BytesIO(raw))
        self.assertEqual(tuple(wb.sheetnames), OIL_GAS_SHEET_ORDER)

    def test_electricity_official_import(self):
        self.assertTrue(FIXTURE.exists(), f"missing fixture {FIXTURE}")
        with FIXTURE.open("rb") as fh:
            result = self.el_handler.import_workbook(
                title=self.el_title,
                user=self.user,
                file_obj=fh,
                dry_run=False,
                sub_main_id=self.sub_main.id,
            )
        self.assertFalse(result["validation"]["blocking"], result)
        self.assertGreater(result["imported_rows"], 0)
        date_info = Info.objects.filter(
            attribute__title=self.el_title,
            attribute__key="report_date",
            value="2026-08-23",
            archived=False,
        )
        self.assertEqual(date_info.count(), 1)
        metrics = (
            Info.objects.filter(
                attribute__title=self.el_title,
                archived=False,
                row_key=date_info.first().row_key,
            )
            .exclude(attribute__key="report_date")
            .count()
        )
        self.assertGreater(metrics, 10)

    def test_oil_gas_official_import(self):
        self.assertTrue(OIL_FIXTURE.exists(), f"missing fixture {OIL_FIXTURE}")
        with OIL_FIXTURE.open("rb") as fh:
            result = self.og_handler.import_workbook(
                title=self.oil_title,
                user=self.user,
                file_obj=fh,
                dry_run=False,
                sub_main_id=self.sub_main.id,
            )
        self.assertFalse(result["validation"]["blocking"], result)
        self.assertGreaterEqual(result["imported_rows"], 1)
        date_info = Info.objects.filter(
            attribute__title=self.oil_title,
            attribute__key="report_date",
            value="2026-09-08",
            archived=False,
        )
        self.assertEqual(date_info.count(), 1)
        metrics = (
            Info.objects.filter(
                attribute__title=self.oil_title,
                archived=False,
                row_key=date_info.first().row_key,
            )
            .exclude(attribute__key="report_date")
            .count()
        )
        self.assertGreaterEqual(metrics, 10)

    def test_template_keys_match_attributes(self):
        raw = self.handler.build_template_bytes(self.oil_title)
        wb = load_workbook(io.BytesIO(raw))
        ws = wb[SHEET_METRICS]
        headers = [c.value for c in ws[1]]
        key_idx = headers.index(COL_KEY) + 1
        keys = {
            str(ws.cell(row=r, column=key_idx).value).strip()
            for r in range(2, ws.max_row + 1)
            if ws.cell(row=r, column=key_idx).value
        }
        attr_keys = {
            a.key for a in importable_attributes(self.oil_title.id) if a.key
        }
        self.assertEqual(keys, attr_keys)

    def test_new_attribute_appears_in_template(self):
        Attribute.objects.create(
            title=self.oil_title,
            label="مؤشر اختبار ديناميكي",
            type="number",
            required=False,
            key="dynamic_test_metric",
            is_measure=True,
            measure_order=999,
        )
        raw = self.handler.build_template_bytes(self.oil_title)
        wb = load_workbook(io.BytesIO(raw))
        ws = wb[SHEET_METRICS]
        headers = [c.value for c in ws[1]]
        key_idx = headers.index(COL_KEY) + 1
        keys = {
            str(ws.cell(row=r, column=key_idx).value).strip()
            for r in range(2, ws.max_row + 1)
            if ws.cell(row=r, column=key_idx).value
        }
        self.assertIn("dynamic_test_metric", keys)

    def test_oil_import_and_reimport_rules(self):
        attrs = importable_attributes(self.oil_title.id)
        values = {"report_date": "2026-09-05"}
        for a in attrs:
            if a.key and a.key != "report_date" and a.type == "number":
                values[a.key] = "1.5"
        buf = self._fill_template(self.oil_title, values)

        result = self.handler.import_workbook(
            title=self.oil_title,
            user=self.user,
            file_obj=buf,
            dry_run=False,
            sub_main_id=self.sub_main.id,
        )
        self.assertFalse(result["validation"]["blocking"], result)
        self.assertEqual(result["imported_rows"], 1)

        buf2 = self._fill_template(self.oil_title, values)
        result2 = self.handler.import_workbook(
            title=self.oil_title,
            user=self.user,
            file_obj=buf2,
            dry_run=False,
            sub_main_id=self.sub_main.id,
        )
        kinds = {i["kind"] for i in result2["validation"]["issues"]}
        self.assertIn("report_date_exists", kinds)
        self.assertGreater(result2["imported_rows"], 0)

        Info.objects.filter(
            attribute__title=self.oil_title,
            attribute__key="report_date",
            value="2026-09-05",
            archived=False,
        ).update(confirmed=Info.ConfirmStatus.ACCEPT)
        date_info = Info.objects.filter(
            attribute__title=self.oil_title,
            attribute__key="report_date",
            value="2026-09-05",
            archived=False,
        ).first()
        Info.objects.filter(row_key=date_info.row_key).update(
            confirmed=Info.ConfirmStatus.ACCEPT
        )

        buf3 = self._fill_template(self.oil_title, values)
        result3 = self.handler.import_workbook(
            title=self.oil_title,
            user=self.user,
            file_obj=buf3,
            dry_run=False,
            sub_main_id=self.sub_main.id,
        )
        self.assertEqual(result3["imported_rows"], 0)
        self.assertTrue(result3["validation"]["blocking"])

        exported = self.handler.build_export_bytes(
            self.oil_title,
            report_date="2026-09-05",
            sub_main_id=self.sub_main.id,
        )
        wb = load_workbook(io.BytesIO(exported))
        self.assertIn(SHEET_METRICS, wb.sheetnames)
