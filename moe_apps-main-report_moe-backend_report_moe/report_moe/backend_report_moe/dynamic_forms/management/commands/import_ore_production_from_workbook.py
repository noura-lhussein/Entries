"""
Load real ore production & contracts figures from the ministry's "وزارة.xlsx"
workbook into the Info form seeded by seed_ore_production_info_form (Title
«إنتاج الخامات والعقود»). One row_key group per (product, plan year).

Expected sheets: "جص 2025", "جص 2026", "رمال كوارتزية 2025", "رمال كوارتزية 2026".
Each product/year combination is hand-mapped below rather than parsed
generically — the column layout differs slightly between 2025 and 2026 sheets
and there are only four sheets, so a small explicit table is clearer than a
generic parser here.

Usage:
  python manage.py import_ore_production_from_workbook --path "C:\\...\\وزارة.xlsx"
"""

from __future__ import annotations

import uuid
from typing import Any

from django.core.management.base import BaseCommand, CommandError
from django.db import transaction
from master_data.models import OreProduct
from openpyxl import load_workbook

from dynamic_forms.models import Attribute, Info, Title

TITLE_NAME = 'إنتاج الخامات والعقود'


def _num(value: Any) -> float | None:
    if value is None or value == '':
        return None
    try:
        return float(value)
    except (TypeError, ValueError):
        return None


class Command(BaseCommand):
    help = 'Import real ore production figures from وزارة.xlsx into report_moe Info.'

    def add_arguments(self, parser):
        parser.add_argument('--path', type=str, required=True, help='Path to وزارة.xlsx')

    def handle(self, *args, **options):
        path = options['path']
        try:
            wb = load_workbook(path, data_only=True)
        except FileNotFoundError as exc:
            raise CommandError(f'Workbook not found: {path}') from exc

        title = Title.objects.filter(name=TITLE_NAME).first()
        if title is None:
            raise CommandError(
                f'Title "{TITLE_NAME}" not found — run seed_ore_production_info_form first.'
            )
        attrs = {a.key: a for a in Attribute.objects.filter(title=title) if a.key}
        required = {'product', 'plan_year', 'annual_plan_tons', 'h1_plan_tons', 'h1_executed_tons', 'contract_count', 'reserve_text'}
        missing = required - set(attrs)
        if missing:
            raise CommandError(f'Title "{TITLE_NAME}" is missing attributes: {", ".join(sorted(missing))}')

        gypsum = OreProduct.objects.filter(name_ar='الجص').first()
        sand_own = OreProduct.objects.filter(name_ar='رمال كوارتزية', production_type='ذاتي').first()
        sand_contracted = OreProduct.objects.filter(name_ar='رمال كوارتزية', production_type='معهّد').first()
        if not gypsum or not sand_own or not sand_contracted:
            raise CommandError('Ore product catalog rows not found — check the geology_oreproduct seed migration.')

        rows_created = 0
        with transaction.atomic():
            # --- جص 2026 ---
            ws = wb['جص 2026']
            r = ws[3]
            rows_created += self._write_row(
                attrs, product_id=gypsum.id, plan_year=2026,
                annual_plan_tons=r[3].value, h1_plan_tons=r[4].value, h1_executed_tons=r[5].value,
                contract_count=r[7].value,
                reserve_text=self._join_reserve(r[8:]),
            )

            # --- جص 2025 --- (no H1 split that year — plan/executed cover the full year)
            ws = wb['جص 2025']
            r = ws[3]
            rows_created += self._write_row(
                attrs, product_id=gypsum.id, plan_year=2025,
                annual_plan_tons=r[3].value, h1_plan_tons=None, h1_executed_tons=r[4].value,
                contract_count=None,
                reserve_text=self._join_reserve(r[6:]),
            )

            # --- رمال كوارتزية 2026 --- (row 3 = ذاتي, row 4 = معهّد)
            ws = wb['رمال كوارتزية 2026']
            own = ws[3]
            rows_created += self._write_row(
                attrs, product_id=sand_own.id, plan_year=2026,
                annual_plan_tons=own[4].value, h1_plan_tons=own[5].value, h1_executed_tons=own[6].value,
                contract_count=own[8].value,
                reserve_text=self._join_reserve(own[9:]),
            )
            contracted = ws[4]
            rows_created += self._write_row(
                attrs, product_id=sand_contracted.id, plan_year=2026,
                annual_plan_tons=contracted[4].value, h1_plan_tons=contracted[5].value,
                h1_executed_tons=contracted[6].value, contract_count=None, reserve_text='',
            )

            # --- رمال كوارتزية 2025 --- (row 3 = ذاتي, row 4 = معهّد; no H1 split)
            ws = wb['رمال كوارتزية 2025']
            own = ws[3]
            rows_created += self._write_row(
                attrs, product_id=sand_own.id, plan_year=2025,
                annual_plan_tons=own[4].value, h1_plan_tons=None, h1_executed_tons=own[5].value,
                contract_count=None, reserve_text=self._join_reserve(own[7:]),
            )
            contracted = ws[4]
            rows_created += self._write_row(
                attrs, product_id=sand_contracted.id, plan_year=2025,
                annual_plan_tons=contracted[4].value, h1_plan_tons=None,
                h1_executed_tons=contracted[5].value, contract_count=None, reserve_text='',
            )

        self.stdout.write(self.style.SUCCESS(f'Done. Created {rows_created} Info row-groups.'))

    def _join_reserve(self, cells) -> str:
        parts = [str(c.value).strip() for c in cells if c.value not in (None, '')]
        return ' / '.join(parts)

    def _write_row(
        self,
        attrs: dict[str, Attribute],
        *,
        product_id: int,
        plan_year: int,
        annual_plan_tons: Any,
        h1_plan_tons: Any,
        h1_executed_tons: Any,
        contract_count: Any,
        reserve_text: str,
    ) -> int:
        row_key = uuid.uuid4()
        Info.objects.create(
            attribute=attrs['product'], value=str(product_id),
            entity_type='ore_product', entity_id=product_id,
            confirmed=Info.ConfirmStatus.ACCEPT, row_key=row_key,
            commit_note='[ore-production-workbook-import]',
        )
        Info.objects.create(
            attribute=attrs['plan_year'], value=str(plan_year),
            confirmed=Info.ConfirmStatus.ACCEPT, row_key=row_key,
            commit_note='[ore-production-workbook-import]',
        )
        for key, raw in (
            ('annual_plan_tons', annual_plan_tons),
            ('h1_plan_tons', h1_plan_tons),
            ('h1_executed_tons', h1_executed_tons),
            ('contract_count', contract_count),
        ):
            num = _num(raw)
            if num is None:
                continue
            Info.objects.create(
                attribute=attrs[key], value=str(num),
                confirmed=Info.ConfirmStatus.ACCEPT, row_key=row_key,
                commit_note='[ore-production-workbook-import]',
            )
        if reserve_text:
            Info.objects.create(
                attribute=attrs['reserve_text'], value=reserve_text,
                confirmed=Info.ConfirmStatus.ACCEPT, row_key=row_key,
                commit_note='[ore-production-workbook-import]',
            )
        self.stdout.write(f'  product={product_id} year={plan_year} row_key={row_key}')
        return 1
