"""
Stub: import electricity report file (PDF/xlsx) → national Info attributes.

PDF→Info mapping must use the seeded national title
«مؤشرات التقرير اليومي للكهرباء». Prefer Form Builder Excel import for
structured titles. This stub validates the file and, when openpyxl can read
xlsx, maps common column headers to national attributes as WAITING (pending).

Usage:
  python manage.py import_electricity_report_file_to_info --source-file report.xlsx \\
      --sub-main-id 15
  python manage.py import_electricity_report_file_to_info --source-file report.xlsx \\
      --sub-main-id 15 --apply
"""

from __future__ import annotations

import uuid
from pathlib import Path

from django.core.management.base import BaseCommand, CommandError
from django.db import transaction

from dynamic_forms.models import Attribute, Info, SubMainSection, Title

TITLE_NATIONAL = 'مؤشرات التقرير اليومي للكهرباء'
NOTE_TAG = '[electricity-file-import-stub]'

# Common xlsx header aliases → attribute key / Arabic label
COLUMN_MAP: dict[str, tuple[str, ...]] = {
    'report_date': ('report_date', 'تاريخ التقرير', 'date', 'التاريخ'),
    'total_generation_mwh_24h': (
        'total_generation_mwh_24h',
        'إجمالي التوليد 24س (MWh)',
        'total_generation',
        'generation_mwh',
    ),
    'peak_generation_mw': ('peak_generation_mw', 'ذروة التوليد', 'peak_mw', 'peak'),
    'net_generation_mwh': ('net_generation_mwh', 'صافي التوليد', 'net_generation'),
    'available_generated_power': ('available_generated_power', 'القدرة المتاحة'),
    'grid_frequency_hz': ('grid_frequency_hz', 'تردد الشبكة', 'frequency'),
}


def _normalize_header(value: object) -> str:
    return str(value or '').strip().lower().replace(' ', '_')


class Command(BaseCommand):
    help = (
        'Stub electricity file→Info importer (dry-run default). '
        'PDF mapping requires seeded national title; use Form Builder Excel for structured titles. '
        'xlsx: maps common columns to national attributes as confirmed=WAITING when --apply.'
    )

    def add_arguments(self, parser):
        parser.add_argument('--source-file', required=True)
        parser.add_argument('--sub-main-id', type=int, required=True)
        parser.add_argument('--apply', action='store_true')
        parser.add_argument('--dry-run', action='store_true')

    def handle(self, *args, **options):
        path = Path(options['source_file'])
        if not path.exists():
            raise CommandError(f'Source file not found: {path}')

        apply = bool(options['apply']) and not bool(options['dry_run'])
        sub = SubMainSection.objects.filter(pk=options['sub_main_id']).first()
        if sub is None:
            raise CommandError(f'SubMainSection {options["sub_main_id"]} not found.')

        title = Title.objects.filter(name=TITLE_NATIONAL).first()
        if title is None:
            raise CommandError(
                f'Missing national title «{TITLE_NATIONAL}». '
                'Run seed_electricity_daily_info_forms. '
                'PDF→Info mapping must use this seeded national title.'
            )

        self.stdout.write(
            f'File OK: {path} ({path.stat().st_size} bytes). '
            f'National title id={title.id}. '
            'Prefer Form Builder Excel import for structured Titles.'
        )

        if path.suffix.lower() not in {'.xlsx', '.xlsm'}:
            self.stdout.write(
                self.style.WARNING(
                    'Non-xlsx file: PDF→Info is not implemented here. '
                    'Use Form Builder Excel or a dedicated PDF extractor.'
                )
            )
            return

        try:
            from openpyxl import load_workbook
        except ImportError:
            self.stdout.write(self.style.WARNING('openpyxl not installed; skipping column map.'))
            return

        wb = load_workbook(path, read_only=True, data_only=True)
        ws = wb.active
        rows_iter = ws.iter_rows(values_only=True)
        try:
            headers = [_normalize_header(h) for h in next(rows_iter)]
        except StopIteration:
            raise CommandError('Empty workbook.') from None

        key_to_attr: dict[str, Attribute] = {}
        for attr in Attribute.objects.filter(title=title):
            for name in (attr.key, attr.label):
                if name:
                    key_to_attr[_normalize_header(name)] = attr

        col_attr: dict[int, Attribute] = {}
        for idx, header in enumerate(headers):
            for attr_key, aliases in COLUMN_MAP.items():
                if header in {_normalize_header(a) for a in aliases} or header == attr_key:
                    attr = key_to_attr.get(attr_key) or key_to_attr.get(_normalize_header(aliases[0]))
                    if attr is None:
                        for alias in aliases:
                            attr = key_to_attr.get(_normalize_header(alias))
                            if attr:
                                break
                    if attr is not None:
                        col_attr[idx] = attr
                    break

        if not col_attr:
            self.stdout.write(self.style.WARNING('No common electricity columns recognized.'))
            return

        created = skipped = scanned = 0
        for row in rows_iter:
            if not row or all(c is None or str(c).strip() == '' for c in row):
                continue
            scanned += 1
            date_val = None
            for idx, attr in col_attr.items():
                if attr.key in ('report_date',) or attr.label in ('تاريخ التقرير', 'report_date'):
                    if idx < len(row) and row[idx] is not None:
                        date_val = str(row[idx])[:10]
            if not date_val:
                date_val = f'row-{scanned}'
            row_key = uuid.uuid5(uuid.NAMESPACE_URL, f'elec-file-import-{path.name}-{date_val}')
            pairs = []
            for idx, attr in col_attr.items():
                if idx >= len(row) or row[idx] is None:
                    continue
                pairs.append((attr, str(row[idx])))
            if not apply:
                created += len(pairs)
                continue
            with transaction.atomic():
                for attr, value in pairs:
                    exists = Info.objects.filter(
                        attribute=attr, sub_main=sub, row_key=row_key
                    ).exists()
                    if exists:
                        skipped += 1
                        continue
                    Info.objects.create(
                        attribute=attr,
                        sub_main=sub,
                        row_key=row_key,
                        value=value,
                        confirmed=Info.ConfirmStatus.WAITING,
                        commit_note=NOTE_TAG,
                    )
                    created += 1

        mode = 'APPLIED' if apply else 'DRY-RUN'
        self.stdout.write(
            self.style.SUCCESS(
                f'{mode}: rows={scanned} mapped_cols={len(col_attr)} '
                f'created≈{created} skipped≈{skipped} confirmed=waiting'
            )
        )
