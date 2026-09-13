"""Dynamic Excel template and import for any title (attributes from DB).

Columns and rows always come from ``importable_attributes`` — do not hardcode
metric key/label lists for Form Builder templates or exports.
"""

from __future__ import annotations

import io
import uuid
from typing import Any, Callable

from django.db import transaction
from openpyxl import Workbook, load_workbook
from openpyxl.styles import Font, PatternFill
from openpyxl.utils import get_column_letter
from openpyxl.worksheet.datavalidation import DataValidation
from rest_framework.exceptions import ValidationError

from ..duplicate_date import check_submit_duplicate_date
from ..excel_utils import cell_value, norm_text
from ..form_schema import attribute_display_label
from ..location_resolution import build_info_records_from_row
from ..models import Info, SubMainSection, Title
from .constants import COL_MAIN_SECTION, COL_REGION, COL_SUB_MAIN
from .exceptions import ImportRowError
from .measures import (
    DEFAULT_TEMPLATE_LAYOUT,
    LayoutNotSupported,
    build_measures_export_bytes,
    build_measures_template_bytes,
    is_measures_workbook,
    read_measures_rows,
)
from .schema import (
    TYPE_LABELS_AR,
    importable_attributes,
    map_headers,
    region_feeds_required_attribute,
    row_values_from_cells,
    template_column_labels,
    validate_and_collect_attribute_values,
)
from .sector_info_writer import RowGroup, write_sector_info
from .submain import resolve_sub_main_for_import


def _empty_result(**extra) -> dict[str, Any]:
    base = {
        "imported_rows": 0,
        "skipped_rows": 0,
        "errors": [],
        "validation": {
            "blocking": False,
            "expected_columns": [],
            "found_columns": [],
            "missing_columns": [],
            "unknown_columns": [],
            "issues": [],
        },
    }
    base.update(extra)
    return base


def _duplicate_date_error(
    row_attrs: list[tuple[int, str]], *, title_id: int, sub_main_id: int, row: int
) -> dict[str, Any] | None:
    """Reuse the guard the manual form uses, so both entry paths answer alike."""
    items = [{"id": attr_id, "value": value} for attr_id, value in row_attrs]
    conflict = check_submit_duplicate_date(
        title_id=title_id, sub_main_id=sub_main_id, items=items
    )
    if not conflict:
        return None
    return {"row": row, "message": conflict["detail"], "code": conflict["code"]}


def _measures_values_by_key(values_by_label: dict[str, str], attrs: list) -> dict[str, str]:
    """Convert label-keyed measures values to Attribute.key keys."""
    by_key: dict[str, str] = {}
    for attr in attrs:
        key = (attr.key or "").strip()
        if not key:
            continue
        if attr.label in values_by_label:
            by_key[key] = values_by_label[attr.label]
    return by_key


def _accepted_values_by_key(
    title: Title, *, report_date: str, sub_main_id: int | None
) -> dict[str, str]:
    """Load accepted Info for a report date into {Attribute.key: value}."""
    date_attrs = [
        a for a in importable_attributes(title.id) if a.is_report_date
    ]
    if not date_attrs:
        raise ValidationError(
            {"report_date": "هذا العنوان لا يحتوي حقل تاريخ تقرير."}
        )
    date_attr_ids = [a.id for a in date_attrs]
    qs = Info.objects.filter(
        attribute_id__in=date_attr_ids,
        archived=False,
        confirmed=Info.ConfirmStatus.ACCEPT,
        value=report_date,
    )
    if sub_main_id is not None:
        qs = qs.filter(sub_main_id=sub_main_id)
    date_info = qs.order_by("-created_at", "-id").first()
    if not date_info or not date_info.row_key:
        raise ValidationError(
            {"report_date": "لا توجد بيانات معتمدة لهذا التاريخ."}
        )
    siblings = Info.objects.filter(
        row_key=date_info.row_key,
        archived=False,
        confirmed=Info.ConfirmStatus.ACCEPT,
    ).select_related("attribute")
    out: dict[str, str] = {}
    for info in siblings:
        key = (info.attribute.key or "").strip()
        if key:
            out[key] = info.value or ""
    if not out:
        raise ValidationError(
            {"report_date": "لا توجد بيانات معتمدة لهذا التاريخ."}
        )
    return out


class DynamicTitleExcelHandler:
    """Single handler: columns and validation come from Title attributes (DB)."""

    def build_template_bytes(
        self,
        title: Title,
        *,
        sub_main_id: int | None = None,
        layout: str = DEFAULT_TEMPLATE_LAYOUT,
    ) -> bytes:
        """Blank import template from Attribute rows (DB SSOT)."""
        if layout == "measures":
            sub_main = None
            if sub_main_id is not None:
                sub_main = SubMainSection.objects.filter(
                    pk=sub_main_id).first()
            return build_measures_template_bytes(title, sub_main=sub_main)

        attrs = importable_attributes(title.id)
        include_sections = sub_main_id is None
        columns = template_column_labels(
            attrs, include_section_columns=include_sections
        )

        wb = Workbook()
        ws = wb.active
        ws.title = "البيانات"
        ws.sheet_view.rightToLeft = True

        header_font = Font(bold=True, color="FFFFFF")
        header_fill = PatternFill("solid", fgColor="7C6B3E")

        for col_idx, label in enumerate(columns, start=1):
            cell = ws.cell(row=1, column=col_idx, value=label)
            cell.font = header_font
            cell.fill = header_fill
            ws.column_dimensions[get_column_letter(col_idx)].width = 22

        last_row = 500
        for col_idx, label in enumerate(columns, start=1):
            attr = next(
                (
                    a
                    for a in attrs
                    if norm_text(attribute_display_label(a)) == norm_text(label)
                    or norm_text(a.label) == norm_text(label)
                ),
                None,
            )
            if not attr or attr.type != "select":
                continue
            options = [
                o.label for o in attr.options.all() if "," not in o.label
            ]
            if not options:
                continue
            col_letter = get_column_letter(col_idx)
            formula = f'"{",".join(options)}"'
            dv = DataValidation(
                type="list", formula1=formula, allow_blank=not attr.required
            )
            dv.error = "اختر قيمة من القائمة"
            dv.errorTitle = "قيمة غير صالحة"
            ws.add_data_validation(dv)
            dv.add(f"{col_letter}2:{col_letter}{last_row}")

        guide = wb.create_sheet("دليل الحقول")
        guide.sheet_view.rightToLeft = True
        guide_headers = ("الحقل", "النوع", "إلزامي", "الخيارات")
        for c, h in enumerate(guide_headers, start=1):
            cell = guide.cell(row=1, column=c, value=h)
            cell.font = header_font
            cell.fill = header_fill
        row = 2
        fixed_guide = (
            list(columns[:3])
            if include_sections
            else [c for c in columns if c == COL_REGION]
        )
        region_required = include_sections or region_feeds_required_attribute(
            attrs
        )
        for label in fixed_guide:
            guide.cell(row=row, column=1, value=label)
            if label in (COL_MAIN_SECTION, COL_SUB_MAIN):
                guide.cell(row=row, column=2, value="يُحدَّد من الواجهة")
                guide.cell(row=row, column=3, value="لا")
            else:
                guide.cell(row=row, column=2, value="موقع")
                guide.cell(
                    row=row,
                    column=3,
                    value="نعم" if region_required else "لا",
                )
            row += 1
        for attr in attrs:
            opts = "، ".join(o.label for o in attr.options.all())
            guide.cell(row=row, column=1, value=attribute_display_label(attr))
            guide.cell(
                row=row,
                column=2,
                value=TYPE_LABELS_AR.get(attr.type, attr.type),
            )
            guide.cell(
                row=row, column=3, value="نعم" if attr.required else "لا"
            )
            guide.cell(row=row, column=4, value=opts)
            row += 1

        buf = io.BytesIO()
        wb.save(buf)
        return buf.getvalue()

    def build_export_bytes(
        self,
        title: Title,
        *,
        report_date: str,
        sub_main_id: int | None = None,
    ) -> bytes:
        """Export accepted Info into the current Attribute-driven measures shape."""
        sub_main = None
        if sub_main_id is not None:
            sub_main = SubMainSection.objects.filter(pk=sub_main_id).first()
        values_by_key = _accepted_values_by_key(
            title, report_date=report_date, sub_main_id=sub_main_id
        )
        try:
            return build_measures_export_bytes(
                title, values_by_key=values_by_key, sub_main=sub_main
            )
        except LayoutNotSupported as exc:
            raise ValidationError({"detail": str(exc)}) from exc

    def _import_measures_workbook(
        self,
        wb,
        *,
        title: Title,
        attrs: list,
        user,
        dry_run: bool,
        log_action: Callable | None,
        sub_main_id: int | None,
    ) -> dict[str, Any]:
        """Import the long layout via attribute-driven ``write_sector_info``."""
        values_by_label, sheet_issues = read_measures_rows(wb, attrs)
        wb.close()

        if not values_by_label:
            return _empty_result(
                errors=[{"row": 1, "message": "لا توجد قيم للاستيراد في الملف"}],
                validation={
                    "blocking": True,
                    "expected_columns": [],
                    "found_columns": [],
                    "missing_columns": [],
                    "unknown_columns": [],
                    "issues": sheet_issues,
                },
                dry_run=dry_run,
            )

        try:
            sub_main = resolve_sub_main_for_import(
                values_by_label, user, 2, sub_main_id=sub_main_id
            )
        except ImportRowError as exc:
            return _empty_result(
                errors=[{"row": exc.row, "message": exc.message}],
                validation={
                    "blocking": True,
                    "expected_columns": [],
                    "found_columns": [],
                    "missing_columns": [],
                    "unknown_columns": [],
                    "issues": sheet_issues
                    + [
                        {
                            "kind": "sheet_unreadable",
                            "column": title.code or title.name,
                            "message": exc.message,
                        }
                    ],
                },
                dry_run=dry_run,
            )

        if not title.code:
            return _empty_result(
                errors=[
                    {
                        "row": 1,
                        "message": (
                            "هذا العنوان بلا رمز (code)؛ "
                            "عيّن Title.code قبل الاستيراد."
                        ),
                    }
                ],
                validation={
                    "blocking": True,
                    "expected_columns": [],
                    "found_columns": [],
                    "missing_columns": [],
                    "unknown_columns": [],
                    "issues": sheet_issues,
                },
                dry_run=dry_run,
            )

        values_by_key = _measures_values_by_key(values_by_label, attrs)
        report_date = (values_by_key.get("report_date") or "").strip()
        title_code = title.code.strip()
        if report_date:
            row_key = uuid.uuid5(
                uuid.NAMESPACE_URL, f"fb-{title_code}-{report_date}"
            )
        else:
            row_key = uuid.uuid4()

        group = RowGroup(
            title_code=title_code,
            entity_type="",
            entity_id=None,
            row_key=row_key,
            values=values_by_key,
        )
        result = write_sector_info(
            groups=[group],
            user=user,
            sub_main=sub_main,
            dry_run=dry_run,
            extra_issues=sheet_issues,
        )
        result.setdefault("validation", {})
        result["validation"]["found_columns"] = sorted(values_by_key.keys())

        if (
            not dry_run
            and result.get("imported_rows")
            and log_action
            and not result.get("validation", {}).get("blocking")
        ):
            log_action(
                details={
                    "title_id": title.id,
                    "title_name": title.name,
                    "imported_rows": result["imported_rows"],
                    "sub_main_id": sub_main_id,
                    "layout": "measures",
                },
            )
        return result

    def import_workbook(
        self,
        *,
        title: Title,
        user,
        file_obj,
        dry_run: bool = False,
        log_action: Callable | None = None,
        sub_main_id: int | None = None,
    ) -> dict[str, Any]:
        attrs = importable_attributes(title.id)
        if not attrs:
            return _empty_result(
                errors=[
                    {
                        "row": 1,
                        "message": "لا توجد حقول قابلة للاستيراد لهذا العنوان",
                    }
                ],
            )

        wb = load_workbook(file_obj, read_only=True, data_only=True)
        if is_measures_workbook(wb):
            return self._import_measures_workbook(
                wb,
                title=title,
                attrs=attrs,
                user=user,
                dry_run=dry_run,
                log_action=log_action,
                sub_main_id=sub_main_id,
            )

        ws = wb.active
        rows_iter = ws.iter_rows(values_only=False)
        try:
            header_cells = next(rows_iter)
        except StopIteration:
            wb.close()
            return _empty_result(
                errors=[{"row": 1, "message": "الملف فارغ"}],
            )

        headers = [cell_value(c) for c in header_cells]
        require_sections = sub_main_id is None
        col_map, validation = map_headers(
            headers,
            attrs,
            title.id,
            require_section_columns=require_sections,
        )

        if validation["blocking"]:
            wb.close()
            return _empty_result(
                errors=[
                    {
                        "row": 1,
                        "message": issue["message"],
                    }
                    for issue in validation["issues"]
                    if issue["kind"] == "missing_column"
                ],
                validation=validation,
                dry_run=dry_run,
            )

        errors: list[dict[str, Any]] = []
        imported = 0
        skipped = 0
        pending_infos: list[Info] = []
        attrs_by_id = {attr.id: attr for attr in attrs}

        for row_idx, row_cells in enumerate(rows_iter, start=2):
            values = row_values_from_cells(col_map, row_cells, cell_value)
            if not any(values.values()):
                skipped += 1
                continue
            try:
                sub_main = resolve_sub_main_for_import(
                    values, user, row_idx, sub_main_id=sub_main_id
                )
                row_attrs = validate_and_collect_attribute_values(
                    values, attrs, col_map, row_idx
                )
                duplicate = _duplicate_date_error(
                    row_attrs,
                    title_id=title.id,
                    sub_main_id=sub_main.id,
                    row=row_idx,
                )
                if duplicate:
                    errors.append(duplicate)
                    skipped += 1
                    continue
                if not dry_run:
                    row_key = uuid.uuid4()
                    pending_infos.extend(
                        build_info_records_from_row(
                            row_attrs,
                            attrs_by_id,
                            sub_main=sub_main,
                            user=user,
                            row_key=row_key,
                        )
                    )
                imported += 1
            except ImportRowError as exc:
                errors.append({"row": exc.row, "message": exc.message})
            except Exception as exc:
                errors.append({"row": row_idx, "message": str(exc)})

        if dry_run:
            wb.close()
            return {
                "imported_rows": imported,
                "skipped_rows": skipped,
                "errors": errors,
                "dry_run": True,
                "validation": validation,
            }

        if errors and imported == 0:
            wb.close()
            return {
                "imported_rows": 0,
                "skipped_rows": skipped,
                "errors": errors,
                "validation": validation,
            }

        with transaction.atomic():
            if pending_infos:
                Info.objects.bulk_create(pending_infos)
                if log_action:
                    log_action(
                        details={
                            "title_id": title.id,
                            "title_name": title.name,
                            "imported_rows": imported,
                            "info_count": len(pending_infos),
                            "sub_main_id": sub_main_id,
                        },
                    )

        wb.close()
        return {
            "imported_rows": imported,
            "skipped_rows": skipped,
            "errors": errors,
            "validation": validation,
        }
