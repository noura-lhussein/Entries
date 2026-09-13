"""
Import drinking water stations into `water_drinkingwaterstation` (schema `moeds`).

Ported from moeds `water/drinking_water_import.py`. report_moe is the single writer for
this table, so the importer lives here; moeds can only read it.

The source workbook has a fixed column layout — the parsing is positional, matching the
original. Do NOT switch to header-name guessing: an earlier stub here did that and
silently dropped coordinates, dates, tei_id, org_unit and address.

Rows are upserted by `station_code`; there is deliberately no --clear. Wiping this table
is what the single-writer work set out to stop.

Usage:
  python manage.py import_drinking_stations_master --source-file stations.xlsx
  python manage.py import_drinking_stations_master --source-file stations.xlsx --apply
"""

from __future__ import annotations

from datetime import date, datetime
from decimal import Decimal, InvalidOperation
from pathlib import Path
from typing import Any

from django.core.management.base import BaseCommand, CommandError
from django.db import transaction

from master_data.models import DrinkingWaterStation


def _parse_date(value: Any) -> date | None:
    if value is None or value == '':
        return None
    if isinstance(value, datetime):
        return value.date()
    if isinstance(value, date):
        return value
    text = str(value).strip()
    if not text:
        return None
    for fmt in ('%Y-%m-%d', '%d/%m/%Y', '%Y/%m/%d'):
        try:
            return datetime.strptime(text, fmt).date()
        except ValueError:
            continue
    return None


def _parse_decimal(value: Any) -> Decimal | None:
    if value is None or value == '':
        return None
    try:
        return Decimal(str(value).strip())
    except (InvalidOperation, ValueError):
        return None


def _clean_text(value: Any) -> str:
    if value is None:
        return ''
    return str(value).strip()


def parse_drinking_water_row(row: tuple[Any, ...]) -> dict[str, Any] | None:
    """Positional parse of one workbook row. Returns None for unusable rows."""
    if len(row) < 15:
        return None

    station_code = _clean_text(row[5])
    name = _clean_text(row[6])
    if not station_code or not name:
        return None

    latitude = _parse_decimal(row[14])
    longitude = _parse_decimal(row[13])
    if latitude is None or longitude is None:
        return None

    return {
        'tei_id': _clean_text(row[0]),
        'org_unit': _clean_text(row[1]),
        'station_code': station_code,
        'name': name,
        'governorate': _clean_text(row[7]),
        'district': _clean_text(row[8]),
        'subdistrict': _clean_text(row[9]),
        'community': _clean_text(row[10]),
        'address': _clean_text(row[11]),
        'enrollment_date': _parse_date(row[3]),
        'incident_date': _parse_date(row[4]),
        'latitude': latitude,
        'longitude': longitude,
    }


class Command(BaseCommand):
    help = 'Upsert drinking water stations from an xlsx workbook.'

    def add_arguments(self, parser):
        parser.add_argument(
            '--source-file',
            default='c:/Users/moe/Documents/moeds/moe-backend/water/data/drinking_water_stations.xlsx'
        )
        parser.add_argument(
            '--apply',
            action='store_true',
            help='Write to the database. Without it the command only reports what it would do.',
        )

    def handle(self, *args, **options):
        path = Path(options['source_file'])
        if not path.exists():
            raise CommandError(f'Source file not found: {path}')
        if path.suffix.lower() not in {'.xlsx', '.xlsm'}:
            raise CommandError(f'Expected an .xlsx/.xlsm workbook, got `{path.suffix}`.')

        try:
            from openpyxl import load_workbook
        except ImportError as exc:
            raise CommandError('openpyxl is required to import drinking water stations.') from exc

        apply = bool(options['apply'])
        existing = DrinkingWaterStation.objects.count()

        workbook = load_workbook(path, data_only=True, read_only=True)
        try:
            worksheet = workbook.active
            rows = worksheet.iter_rows(values_only=True)
            next(rows, None)  # header

            parsed_rows = []
            skipped = 0
            for row in rows:
                parsed = parse_drinking_water_row(row)
                if parsed is None:
                    skipped += 1
                    continue
                parsed_rows.append(parsed)
        finally:
            workbook.close()

        if not apply:
            codes = {r['station_code'] for r in parsed_rows}
            known = set(
                DrinkingWaterStation.objects
                .filter(station_code__in=codes)
                .values_list('station_code', flat=True)
            )
            self.stdout.write(
                self.style.SUCCESS(
                    f'DRY-RUN: {len(parsed_rows)} usable rows ({skipped} skipped). '
                    f'Would create {len(codes - known)}, update {len(known)}. '
                    f'Current table count={existing}. Re-run with --apply.'
                )
            )
            return

        created = updated = 0
        with transaction.atomic():
            for parsed in parsed_rows:
                _, was_created = DrinkingWaterStation.objects.update_or_create(
                    station_code=parsed['station_code'],
                    defaults=parsed,
                )
                if was_created:
                    created += 1
                else:
                    updated += 1

        self.stdout.write(
            self.style.SUCCESS(
                f'APPLIED: created={created} updated={updated} skipped={skipped} '
                f'(table count {existing} → {DrinkingWaterStation.objects.count()})'
            )
        )
