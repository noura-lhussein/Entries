from __future__ import annotations

import re
from datetime import date, datetime
from pathlib import Path
from typing import Any

from django.db import transaction

from .drinking_water_survey_schema import (
    SURVEY_UPDATE_FIELDS,
    build_portal_column_map,
    is_portal_survey_header,
)
from .models import DrinkingWaterStation

try:
    from openpyxl import load_workbook
except ImportError:
    load_workbook = None

NON_OPERATIONAL_REASON_MAP = {
    'other reasons': 'other',
    'other': 'other',
    'due to theft and vandalism': 'theft_vandalism',
    'theft and vandalism': 'theft_vandalism',
    'theft_vandalism': 'theft_vandalism',
    'completely destroyed': 'completely_destroyed',
    'completely_destroyed': 'completely_destroyed',
    'requires emergency maintenance': 'emergency_maintenance',
    'emergency maintenance': 'emergency_maintenance',
    'emergency_maintenance': 'emergency_maintenance',
    'administrative reasons': 'administrative',
    'administrative': 'administrative',
    'requires routine maintenance': 'routine_maintenance',
    'routine maintenance': 'routine_maintenance',
    'routine_maintenance': 'routine_maintenance',
    'ongoing maintenance': 'ongoing_maintenance',
    'ongoing_maintenance': 'ongoing_maintenance',
}


def _clean_text(value: Any) -> str:
    if value is None:
        return ''
    return str(value).strip()


def _parse_date(value: Any) -> date | None:
    if value is None or value == '':
        return None
    if isinstance(value, datetime):
        return value.date()
    if isinstance(value, date):
        return value
    text = str(value).strip()
    if not text:
        return None
    for fmt in ('%Y-%m-%d', '%d/%m/%Y', '%Y/%m/%d'):
        try:
            return datetime.strptime(text, fmt).date()
        except ValueError:
            continue
    return None


def _parse_yes_no(value: Any) -> bool | None:
    text = _clean_text(value).lower()
    if not text:
        return None
    if text in {'yes', 'نعم', 'y'}:
        return True
    if text in {'no', 'لا', 'n'}:
        return False
    return None


def _parse_level_prefix(value: Any) -> str:
    text = _clean_text(value)
    if not text:
        return ''
    match = re.match(r'^([123])\b', text)
    if match:
        return match.group(1)
    if text in {'1', '2', '3'}:
        return text
    return ''


def _parse_building_condition(value: Any) -> str:
    return _parse_level_prefix(value)


def _parse_safety_procedures(value: Any) -> str:
    text = _clean_text(value).lower()
    if not text:
        return ''
    if text == 'partially':
        return 'partially'
    if text == 'yes':
        return 'yes'
    if text == 'no':
        return 'no'
    return ''


def _parse_rehabilitation_type(value: Any) -> str:
    text = _clean_text(value).lower()
    if 'complete' in text:
        return 'complete'
    if 'partial' in text:
        return 'partial'
    return ''


def _parse_lab_equipment_status(value: Any) -> str:
    return _parse_safety_procedures(value)


def _parse_non_operational_reason(value: Any) -> str:
    text = _clean_text(value).lower()
    if not text:
        return ''
    return NON_OPERATIONAL_REASON_MAP.get(text, 'other')


def _parse_productivity(value: Any) -> str:
    text = _clean_text(value)
    if not text:
        return ''
    if text.isdigit():
        num = int(text)
        if num >= 100:
            return '100'
        if num >= 76:
            return '76_99'
        if num >= 51:
            return '51_75'
        if num >= 26:
            return '26_50'
        if num >= 1:
            return '1_25'
        return '0'
    normalized = text.replace('-', '_').replace(' ', '').lower()
    if normalized in {'0', '1_25', '26_50', '51_75', '76_99', '100'}:
        return normalized
    return normalized


def parse_survey_fields(data: dict[str, Any]) -> dict[str, Any] | None:
    tei_id = _clean_text(data.get('tei_id'))
    if not tei_id:
        return None

    has_public_grid_supply = _parse_yes_no(data.get('has_public_grid_supply'))
    needs_solar_installation = _parse_yes_no(data.get('needs_solar_installation'))
    solar_power_available = _parse_yes_no(data.get('solar_power_available'))
    generator_available = _parse_yes_no(data.get('generator_available'))

    return {
        'tei_id': tei_id,
        'enrollment_date': _parse_date(data.get('enrollment_date')),
        'org_unit': _clean_text(data.get('org_unit')),
        'building_condition': _parse_building_condition(data.get('building_condition')),
        'is_operational': _parse_yes_no(data.get('is_operational')),
        'non_operational_reason': _parse_non_operational_reason(data.get('non_operational_reason')),
        'previously_rehabilitated': _parse_yes_no(data.get('previously_rehabilitated')),
        'rehabilitation_type': _parse_rehabilitation_type(data.get('rehabilitation_type')),
        'has_water_hammer_protection': _parse_yes_no(data.get('has_water_hammer_protection')),
        'water_hammer_efficiency': _parse_level_prefix(data.get('water_hammer_efficiency')),
        'needs_solar_installation': needs_solar_installation,
        'needs_solar_power': needs_solar_installation,
        'solar_space_available': _parse_yes_no(data.get('solar_space_available')),
        'safety_procedures': _parse_safety_procedures(data.get('safety_procedures')),
        'has_water_tanks': _parse_yes_no(data.get('has_water_tanks')),
        'has_public_grid_supply': has_public_grid_supply,
        'has_grid_power': has_public_grid_supply,
        'grid_connection_working': _parse_yes_no(data.get('grid_connection_working')),
        'electrical_connection_efficiency': _parse_level_prefix(data.get('electrical_connection_efficiency')),
        'transformer_efficiency': _parse_level_prefix(data.get('transformer_efficiency')),
        'electrical_panel_efficiency': _parse_level_prefix(data.get('electrical_panel_efficiency')),
        'grid_power_productivity': _parse_productivity(data.get('grid_power_productivity')),
        'solar_power_available': solar_power_available,
        'solar_power_productivity': _parse_productivity(data.get('solar_power_productivity')),
        'solar_system_efficiency': _parse_level_prefix(data.get('solar_system_efficiency')),
        'generator_available': generator_available,
        'alternative_power_source': (
            True if solar_power_available or generator_available else
            False if solar_power_available is False and generator_available is False else
            _parse_yes_no(data.get('alternative_power_source'))
        ),
        'is_well_station': _parse_yes_no(data.get('is_well_station')),
        'is_filtration_station': _parse_yes_no(data.get('is_filtration_station')),
        'is_boosting_station': _parse_yes_no(data.get('is_boosting_station')),
        'is_water_analyzed': _parse_yes_no(data.get('is_water_analyzed')),
        'lab_equipment_status': _parse_lab_equipment_status(data.get('lab_equipment_status')),
    }


def parse_survey_row(row: tuple[Any, ...]) -> dict[str, Any] | None:
    if len(row) < 81:
        return None
    return parse_survey_fields({
        'tei_id': row[1],
        'enrollment_date': row[3],
        'org_unit': row[4],
        'building_condition': row[5],
        'is_operational': row[6],
        'non_operational_reason': row[7],
        'previously_rehabilitated': row[16],
        'rehabilitation_type': row[17],
        'has_water_hammer_protection': row[19],
        'water_hammer_efficiency': row[20],
        'needs_solar_installation': row[22],
        'solar_space_available': row[23],
        'safety_procedures': row[25],
        'has_water_tanks': row[27],
        'has_public_grid_supply': row[31],
        'grid_connection_working': row[32],
        'electrical_connection_efficiency': row[33],
        'transformer_efficiency': row[34],
        'electrical_panel_efficiency': row[35],
        'grid_power_productivity': row[36],
        'solar_power_available': row[40],
        'solar_power_productivity': row[41],
        'solar_system_efficiency': row[47],
        'generator_available': row[48],
        'is_well_station': row[54],
        'is_filtration_station': row[56],
        'is_boosting_station': row[70],
        'is_water_analyzed': row[76],
        'lab_equipment_status': row[80],
    })


def parse_portal_survey_row(row: tuple[Any, ...], column_map: dict[str, int]) -> dict[str, Any] | None:
    def cell(field: str) -> Any:
        index = column_map.get(field)
        if index is None or index >= len(row):
            return None
        return row[index]

    data = {field: cell(field) for field in ('tei_id', *SURVEY_UPDATE_FIELDS)}
    return parse_survey_fields(data)


def import_drinking_water_survey_workbook(path: Path) -> dict[str, int]:
    if load_workbook is None:
        raise RuntimeError('openpyxl is required to import drinking water survey data')

    workbook = load_workbook(path, data_only=True, read_only=True)
    worksheet = workbook.active
    rows = list(worksheet.iter_rows(values_only=True))
    workbook.close()

    if not rows:
        return {'updated': 0, 'missing': 0, 'skipped': 0, 'format': 'empty'}

    portal_format = is_portal_survey_header(rows[0])
    column_map = build_portal_column_map(rows[0]) if portal_format else {}
    data_rows = rows[1:]

    updated = 0
    missing = 0
    skipped = 0

    with transaction.atomic():
        for row in data_rows:
            if portal_format:
                parsed = parse_portal_survey_row(row, column_map)
            else:
                parsed = parse_survey_row(row)

            if not parsed:
                skipped += 1
                continue

            tei_id = parsed.pop('tei_id')
            station = DrinkingWaterStation.objects.filter(tei_id=tei_id).first()
            if not station:
                missing += 1
                continue

            update_fields = list(parsed.keys())
            for field in update_fields:
                setattr(station, field, parsed[field])
            station.save(update_fields=update_fields)
            updated += 1

    return {
        'updated': updated,
        'missing': missing,
        'skipped': skipped,
        'format': 'portal' if portal_format else 'tei',
    }
