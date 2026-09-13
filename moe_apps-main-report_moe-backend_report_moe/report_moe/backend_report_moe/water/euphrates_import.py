from __future__ import annotations

import re
import unicodedata
from datetime import date, datetime, timedelta
from decimal import Decimal, InvalidOperation
from pathlib import Path
from typing import Any, Iterable

from django.db import transaction

from .models import EuphratesCascadeReading

try:
    from openpyxl import load_workbook
except ImportError:
    load_workbook = None

try:
    import xlrd
except ImportError:
    xlrd = None


def _normalize_filename(value: str) -> str:
    return unicodedata.normalize('NFKC', value)


def _parse_decimal(value: Any) -> Decimal | None:
    if value is None or value == '':
        return None
    try:
        return Decimal(str(value).strip().replace(',', ''))
    except (InvalidOperation, ValueError):
        return None


_MIN_READING_YEAR = 1998
_MAX_READING_YEAR = 2030


def _parse_date(value: Any) -> date | None:
    if value is None or value == '':
        return None
    if isinstance(value, datetime):
        return value.date()
    if isinstance(value, date):
        return value
    if isinstance(value, (int, float)) and 25000 < value < 60000:
        return date(1899, 12, 30) + timedelta(days=int(value))
    text = str(value).strip()
    if not text:
        return None
    for fmt in ('%Y-%m-%d %H:%M:%S', '%Y-%m-%d', '%d/%m/%Y', '%d-%m-%Y'):
        try:
            return datetime.strptime(text, fmt).date()
        except ValueError:
            continue
    return None


def _years_from_filename(file_path: Path) -> set[int] | None:
    name = _normalize_filename(file_path.stem)
    match = re.search(r'(\d{4})\s*[-–]\s*(\d{4})', name)
    if match:
        start, end = int(match.group(1)), int(match.group(2))
        return set(range(start, end + 1))
    match = re.search(r'(\d{4})', name)
    if match:
        year = int(match.group(1))
        return {year - 1, year, year + 1}
    return None


def _is_plausible_reading_date(reading_date: date, file_path: Path) -> bool:
    if reading_date.year < _MIN_READING_YEAR or reading_date.year > _MAX_READING_YEAR:
        return False
    allowed_years = _years_from_filename(file_path)
    if allowed_years is None:
        return True
    return reading_date.year in allowed_years


def _normalize_header(value: Any) -> str:
    text = str(value or '').strip()
    text = re.sub(r'\s+', ' ', text)
    return text.replace('سدال', 'سد ').replace('بحيرة الأسد', 'الأسد')


def _header_matches(normalized: str, *parts: str) -> bool:
    return all(part in normalized for part in parts)


def _header_excludes(normalized: str, *parts: str) -> bool:
    return not any(part in normalized for part in parts)


def _is_furat_column(normalized: str) -> bool:
    return 'فرات' in normalized or 'الأسد' in normalized


def _build_column_map(headers: list[Any]) -> dict[str, int | None]:
    mapping: dict[str, int | None] = {
        'inflow_jarabulus': None,
        'tishreen_level_m': None,
        'tishreen_storage_mcm': None,
        'tishreen_outflow': None,
        'tishreen_generation_mwh': None,
        'furat_level_m': None,
        'furat_storage_mcm': None,
        'furat_outflow': None,
        'furat_generation_mwh': None,
        'kadiran_outflow': None,
        'kadiran_generation_mwh': None,
        'al_jalab_discharge': None,
        'total_generation_mwh': None,
    }

    for index, header in enumerate(headers):
        normalized = _normalize_header(header)
        if not normalized or normalized == 'التاريخ':
            continue

        if mapping['inflow_jarabulus'] is None and _header_matches(normalized, 'وارد', 'جرابلس'):
            mapping['inflow_jarabulus'] = index
            continue

        if mapping['tishreen_level_m'] is None and _header_matches(normalized, 'منسوب', 'تشرين'):
            mapping['tishreen_level_m'] = index
            continue

        if mapping['tishreen_storage_mcm'] is None and _header_matches(normalized, 'حجم', 'تخزين', 'تشرين'):
            mapping['tishreen_storage_mcm'] = index
            continue

        if (
            mapping['tishreen_outflow'] is None
            and _header_matches(normalized, 'ممرر', 'تشرين')
            and _header_excludes(normalized, 'بعث', 'كديران', 'بوكمال')
        ):
            mapping['tishreen_outflow'] = index
            continue

        if (
            mapping['tishreen_generation_mwh'] is None
            and _header_matches(normalized, 'مولدة', 'تشرين')
            and _header_excludes(normalized, 'مستجرة', 'ثلاث')
        ):
            mapping['tishreen_generation_mwh'] = index
            continue

        if mapping['furat_level_m'] is None and _header_matches(normalized, 'منسوب') and _is_furat_column(normalized):
            mapping['furat_level_m'] = index
            continue

        if mapping['furat_storage_mcm'] is None and _header_matches(normalized, 'حجم', 'تخزين') and _is_furat_column(normalized):
            mapping['furat_storage_mcm'] = index
            continue

        if (
            mapping['furat_outflow'] is None
            and _header_matches(normalized, 'ممرر')
            and _is_furat_column(normalized)
            and _header_excludes(normalized, 'بوكمال', 'بعث', 'كديران')
        ):
            mapping['furat_outflow'] = index
            continue

        if (
            mapping['furat_generation_mwh'] is None
            and _header_matches(normalized, 'مولدة')
            and _is_furat_column(normalized)
            and _header_excludes(normalized, 'مستجرة', 'ثلاث', 'تشرين', 'بعث', 'كديران')
        ):
            mapping['furat_generation_mwh'] = index
            continue

        if (
            mapping['kadiran_outflow'] is None
            and _header_matches(normalized, 'ممرر')
            and ('بعث' in normalized or 'كديران' in normalized)
        ):
            mapping['kadiran_outflow'] = index
            continue

        if (
            mapping['kadiran_generation_mwh'] is None
            and _header_matches(normalized, 'مولدة')
            and ('بعث' in normalized or 'كديران' in normalized)
        ):
            mapping['kadiran_generation_mwh'] = index
            continue

        if mapping['al_jalab_discharge'] is None and 'جلاب' in normalized:
            mapping['al_jalab_discharge'] = index
            continue

        if mapping['total_generation_mwh'] is None and _header_matches(normalized, 'ثلاث'):
            mapping['total_generation_mwh'] = index
            continue

    return mapping


def _cell_value(row: tuple[Any, ...] | list[Any], index: int | None) -> Any:
    if index is None or index >= len(row):
        return None
    return row[index]


def _reading_from_row(
    row: tuple[Any, ...] | list[Any],
    *,
    col_map: dict[str, int | None],
    reading_date: date,
    report_label: str,
    source_file: str,
) -> EuphratesCascadeReading:
    return EuphratesCascadeReading(
        reading_date=reading_date,
        inflow_jarabulus=_parse_decimal(_cell_value(row, col_map['inflow_jarabulus'])),
        tishreen_level_m=_parse_decimal(_cell_value(row, col_map['tishreen_level_m'])),
        tishreen_storage_mcm=_parse_decimal(_cell_value(row, col_map['tishreen_storage_mcm'])),
        tishreen_outflow=_parse_decimal(_cell_value(row, col_map['tishreen_outflow'])),
        tishreen_generation_mwh=_parse_decimal(_cell_value(row, col_map['tishreen_generation_mwh'])),
        furat_level_m=_parse_decimal(_cell_value(row, col_map['furat_level_m'])),
        furat_storage_mcm=_parse_decimal(_cell_value(row, col_map['furat_storage_mcm'])),
        furat_outflow=_parse_decimal(_cell_value(row, col_map['furat_outflow'])),
        furat_generation_mwh=_parse_decimal(_cell_value(row, col_map['furat_generation_mwh'])),
        kadiran_outflow=_parse_decimal(_cell_value(row, col_map['kadiran_outflow'])),
        kadiran_generation_mwh=_parse_decimal(_cell_value(row, col_map['kadiran_generation_mwh'])),
        al_jalab_discharge=_parse_decimal(_cell_value(row, col_map['al_jalab_discharge'])),
        total_generation_mwh=_parse_decimal(_cell_value(row, col_map['total_generation_mwh'])),
        report_label=report_label,
        source_file=source_file,
    )


def _find_header_row(rows: list[tuple[Any, ...]]) -> tuple[int, list[Any]] | None:
    for index, row in enumerate(rows[:8]):
        if row and _normalize_header(row[0]) == 'التاريخ':
            return index, list(row)
    return None


def _is_skipped_sheet(name: str) -> bool:
    lowered = name.strip().lower()
    if lowered.startswith('ورقة') or lowered in {'sheet1', 'sheet2', 'sheet3'}:
        return True
    skip_markers = ('خطة', 'ممرر البعث', 'اصل البوكمال')
    return any(marker in name for marker in skip_markers)


def _sheet_names_to_import(sheet_names: Iterable[str]) -> list[str]:
    return [name for name in sheet_names if not _is_skipped_sheet(name)]


def _bulk_upsert(batch: list[EuphratesCascadeReading]) -> None:
    if not batch:
        return
    deduped: dict[date, EuphratesCascadeReading] = {}
    for row in batch:
        deduped[row.reading_date] = row
    rows = list(deduped.values())
    with transaction.atomic():
        EuphratesCascadeReading.objects.bulk_create(
            rows,
            update_conflicts=True,
            unique_fields=['reading_date'],
            update_fields=[
                'inflow_jarabulus', 'tishreen_level_m', 'tishreen_storage_mcm',
                'tishreen_outflow', 'tishreen_generation_mwh', 'furat_level_m',
                'furat_storage_mcm', 'furat_outflow', 'furat_generation_mwh',
                'kadiran_outflow', 'kadiran_generation_mwh', 'al_jalab_discharge',
                'total_generation_mwh', 'report_label', 'source_file',
            ],
        )


def _is_water_info_filename(name: str) -> bool:
    normalized = _normalize_filename(name)
    return 'المعلومات المائية' in normalized or 'معلومات مائية' in normalized


def is_euphrates_workbook(path: Path) -> bool:
    if path.name.startswith('~$'):
        return False
    if path.suffix.lower() not in {'.xlsx', '.xls'}:
        return False
    name = _normalize_filename(path.name)
    lowered = name.lower()
    return (
        'سدود الفرات' in name
        or 'euphrates' in lowered
        or _is_water_info_filename(name)
    )


def _import_legacy_monthly_sheet(rows: list[tuple[Any, ...]], file_path: Path) -> int:
    if len(rows) < 3:
        return 0
    report_label = ' '.join(str(cell) for cell in rows[0] if cell).strip()
    header = rows[1]
    if _normalize_header(header[0]) != 'التاريخ':
        return 0

    col_map = _build_column_map(list(header))
    batch: list[EuphratesCascadeReading] = []
    processed = 0
    for row in rows[2:]:
        if not row or row[0] in (None, ''):
            continue
        reading_date = _parse_date(row[0])
        if not reading_date or not _is_plausible_reading_date(reading_date, file_path):
            continue
        batch.append(
            _reading_from_row(
                row,
                col_map=col_map,
                reading_date=reading_date,
                report_label=report_label,
                source_file=file_path.name,
            ),
        )
        processed += 1
    _bulk_upsert(batch)
    return processed


def _import_water_info_sheet(
    rows: list[tuple[Any, ...]],
    *,
    sheet_name: str,
    file_path: Path,
) -> int:
    header_info = _find_header_row(rows)
    if not header_info:
        return 0

    header_index, header = header_info
    col_map = _build_column_map(header)
    report_label = ''
    if header_index > 0:
        report_label = ' '.join(str(cell) for cell in rows[header_index - 1] if cell).strip()
    if not report_label:
        report_label = f'{file_path.stem} — {sheet_name}'

    batch: list[EuphratesCascadeReading] = []
    processed = 0
    for row in rows[header_index + 1:]:
        if not row or row[0] in (None, ''):
            continue
        reading_date = _parse_date(row[0])
        if not reading_date or not _is_plausible_reading_date(reading_date, file_path):
            continue
        batch.append(
            _reading_from_row(
                row,
                col_map=col_map,
                reading_date=reading_date,
                report_label=report_label,
                source_file=file_path.name,
            ),
        )
        processed += 1
    _bulk_upsert(batch)
    return processed


def _xls_rows(sheet, workbook) -> list[tuple[Any, ...]]:
    rows: list[tuple[Any, ...]] = []
    for row_index in range(sheet.nrows):
        values: list[Any] = []
        for col_index in range(sheet.ncols):
            cell_type = sheet.cell_type(row_index, col_index)
            value = sheet.cell_value(row_index, col_index)
            if cell_type == xlrd.XL_CELL_DATE:
                value = xlrd.xldate.xldate_as_datetime(value, workbook.datemode).date()
            values.append(value)
        rows.append(tuple(values))
    return rows


def import_euphrates_workbook(file_path: Path) -> int:
    suffix = file_path.suffix.lower()
    if suffix == '.xlsx':
        if load_workbook is None:
            raise RuntimeError('openpyxl is required to import .xlsx Euphrates workbooks')
        workbook = load_workbook(str(file_path), read_only=True, data_only=True)
        total = 0
        if _is_water_info_filename(file_path.name) or file_path.suffix.lower() == '.xls':
            sheet_names = _sheet_names_to_import(workbook.sheetnames)
            for sheet_name in sheet_names:
                rows = list(workbook[sheet_name].iter_rows(values_only=True))
                total += _import_water_info_sheet(rows, sheet_name=sheet_name, file_path=file_path)
        else:
            sheet = workbook.active
            rows = list(sheet.iter_rows(values_only=True))
            total += _import_legacy_monthly_sheet(rows, file_path)
        workbook.close()
        return total

    if suffix == '.xls':
        if xlrd is None:
            raise RuntimeError('xlrd is required to import .xls Euphrates workbooks')
        workbook = xlrd.open_workbook(str(file_path))
        total = 0
        sheet_names = _sheet_names_to_import(workbook.sheet_names())
        for sheet_name in sheet_names:
            sheet = workbook.sheet_by_name(sheet_name)
            rows = _xls_rows(sheet, workbook)
            total += _import_water_info_sheet(rows, sheet_name=sheet_name, file_path=file_path)
        return total

    return 0


def import_euphrates_folder(root: Path, *, clear: bool = False) -> dict[str, int]:
    if not root.is_dir():
        raise FileNotFoundError(f'Dam data folder not found: {root}')

    files = sorted(
        path
        for path in root.rglob('*')
        if path.is_file() and is_euphrates_workbook(path)
    )
    if not files:
        raise FileNotFoundError(f'No Euphrates cascade workbooks found under {root}')

    if clear:
        EuphratesCascadeReading.objects.all().delete()

    total = 0
    for file_path in files:
        total += import_euphrates_workbook(file_path)
    return {'files': len(files), 'rows': total}
