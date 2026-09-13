from __future__ import annotations

from datetime import date, datetime
from decimal import Decimal, InvalidOperation
from pathlib import Path
from typing import Any

from django.db import transaction

from .models import DrinkingWaterStation

try:
    from openpyxl import load_workbook
except ImportError:
    load_workbook = None


def _parse_date(value: Any) -> date | None:
    if value is None or value == '':
        return None
    if isinstance(value, datetime):
        return value.date()
    if isinstance(value, date):
        return value
    text = str(value).strip()
    if not text:
        return None
    for fmt in ('%Y-%m-%d', '%d/%m/%Y', '%Y/%m/%d'):
        try:
            return datetime.strptime(text, fmt).date()
        except ValueError:
            continue
    return None


def _parse_decimal(value: Any) -> Decimal | None:
    if value is None or value == '':
        return None
    try:
        return Decimal(str(value).strip())
    except (InvalidOperation, ValueError):
        return None


def _clean_text(value: Any) -> str:
    if value is None:
        return ''
    return str(value).strip()


def parse_drinking_water_row(row: tuple[Any, ...]) -> dict[str, Any] | None:
    if len(row) < 15:
        return None

    station_code = _clean_text(row[5])
    name = _clean_text(row[6])
    if not station_code or not name:
        return None

    latitude = _parse_decimal(row[14])
    longitude = _parse_decimal(row[13])
    if latitude is None or longitude is None:
        return None

    return {
        'tei_id': _clean_text(row[0]),
        'org_unit': _clean_text(row[1]),
        'station_code': station_code,
        'name': name,
        'governorate': _clean_text(row[7]),
        'district': _clean_text(row[8]),
        'subdistrict': _clean_text(row[9]),
        'community': _clean_text(row[10]),
        'address': _clean_text(row[11]),
        'enrollment_date': _parse_date(row[3]),
        'incident_date': _parse_date(row[4]),
        'latitude': latitude,
        'longitude': longitude,
    }


def import_drinking_water_workbook(path: Path, *, clear: bool = False) -> dict[str, int]:
    if load_workbook is None:
        raise RuntimeError('openpyxl is required to import drinking water stations')

    workbook = load_workbook(path, data_only=True, read_only=True)
    worksheet = workbook.active
    rows = worksheet.iter_rows(values_only=True)
    next(rows, None)

    created = 0
    updated = 0
    skipped = 0

    with transaction.atomic():
        if clear:
            DrinkingWaterStation.objects.all().delete()

        for row in rows:
            parsed = parse_drinking_water_row(row)
            if not parsed:
                skipped += 1
                continue

            _, was_created = DrinkingWaterStation.objects.update_or_create(
                station_code=parsed['station_code'],
                defaults=parsed,
            )
            if was_created:
                created += 1
            else:
                updated += 1

    workbook.close()
    return {'created': created, 'updated': updated, 'skipped': skipped}
