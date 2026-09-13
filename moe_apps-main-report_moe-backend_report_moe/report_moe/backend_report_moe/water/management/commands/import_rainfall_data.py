from __future__ import annotations

import os
from datetime import date, datetime
from decimal import Decimal, InvalidOperation
from pathlib import Path

from django.core.management.base import BaseCommand, CommandError
from django.db import transaction

from water.basin_catalog import BASIN_SPECS, BasinSpec
from water.coordinates import utm_to_wgs84
from water.models import RainfallBasin, RainfallObservation, RainfallStation

try:
    import xlrd
except ImportError as exc:
    xlrd = None
    _XLRD_IMPORT_ERROR = exc
else:
    _XLRD_IMPORT_ERROR = None

try:
    import openpyxl
except ImportError as exc:
    openpyxl = None
    _OPENPYXL_IMPORT_ERROR = exc
else:
    _OPENPYXL_IMPORT_ERROR = None


def _resolve_basin(folder_name: str) -> BasinSpec | None:
    for spec in BASIN_SPECS:
        if spec.folder_key in folder_name:
            return spec
    return None


def _parse_date(value, *, xls_datemode: int | None = None) -> date | None:
    if value is None or value == '':
        return None
    if isinstance(value, datetime):
        return value.date()
    if isinstance(value, date):
        return value
    if isinstance(value, float) and xlrd is not None and xls_datemode is not None:
        try:
            return xlrd.xldate_as_datetime(value, xls_datemode).date()
        except Exception:
            return None
    text = str(value).strip()
    if not text:
        return None
    for fmt in ('%Y-%m-%d', '%d/%m/%Y', '%Y/%m/%d'):
        try:
            return datetime.strptime(text, fmt).date()
        except ValueError:
            continue
    return None


def _parse_decimal(value) -> Decimal:
    if value is None or value == '':
        return Decimal('0')
    try:
        return Decimal(str(value).strip())
    except (InvalidOperation, ValueError):
        return Decimal('0')


def _collect_workbook_files(root: Path) -> list[Path]:
    files: list[Path] = []
    for path in root.rglob('*'):
        if not path.is_file():
            continue
        if path.name.startswith('~$'):
            continue
        if path.suffix.lower() in {'.xls', '.xlsx'}:
            files.append(path)
    return sorted(files)


def _folder_diagnostics(root: Path, limit: int = 40) -> str:
    lines = [f'Contents under {root}:']
    try:
        entries = sorted(root.iterdir(), key=lambda p: p.name)
    except OSError as exc:
        return f'Could not list folder: {exc}'
    if not entries:
        return f'Folder is empty: {root}'
    for entry in entries[:limit]:
        kind = 'dir' if entry.is_dir() else 'file'
        lines.append(f'  [{kind}] {entry.name}')
    if len(entries) > limit:
        lines.append(f'  … and {len(entries) - limit} more')
    exts: dict[str, int] = {}
    for path in root.rglob('*'):
        if path.is_file():
            key = path.suffix.lower() or '(no extension)'
            exts[key] = exts.get(key, 0) + 1
    if exts:
        lines.append('File extensions: ' + ', '.join(f'{k}={v}' for k, v in sorted(exts.items())))
    return '\n'.join(lines)


class Command(BaseCommand):
    help = 'Import agricultural rainfall station XLS/XLSX files into the water database.'

    def add_arguments(self, parser):
        parser.add_argument(
            '--path',
            default=os.getenv(
                'RAINFALL_DATA_DIR',
                str(Path(__file__).resolve().parents[4] / 'هطول محطات الزراعة'),
            ),
            help='Root folder containing basin subfolders with .xls/.xlsx files.',
        )
        parser.add_argument(
            '--clear',
            action='store_true',
            help='Delete existing rainfall observations before import.',
        )

    def handle(self, *args, **options):
        root = Path(options['path']).expanduser().resolve()
        if not root.is_dir():
            raise CommandError(f'Rainfall data folder not found: {root}')

        for spec in BASIN_SPECS:
            RainfallBasin.objects.update_or_create(
                slug=spec.slug,
                defaults={'name_en': spec.name_en, 'name_ar': spec.name_ar},
            )

        if options['clear']:
            self.stdout.write('Clearing existing rainfall observations…')
            RainfallObservation.objects.all().delete()

        files = _collect_workbook_files(root)
        if not files:
            raise CommandError(
                'No .xls/.xlsx files found under '
                f'{root}\n\n{_folder_diagnostics(root)}\n\n'
                'Expected basin folders named like: البادية، الخابور، الساحل، العاصي، '
                'الفرات، اليرموك، بردى — each containing station workbooks.'
            )

        total_rows = 0
        skipped = 0
        for file_path in files:
            basin_spec = _resolve_basin(file_path.parent.name) or _resolve_basin(file_path.name)
            if not basin_spec:
                skipped += 1
                self.stdout.write(self.style.WARNING(f'Skip (unknown basin): {file_path}'))
                continue
            basin = RainfallBasin.objects.get(slug=basin_spec.slug)
            rows = self._import_file(file_path, basin)
            total_rows += rows
            self.stdout.write(f'Imported {rows:,} rows ({basin_spec.slug}) ← {file_path.name}')

        if total_rows == 0 and skipped:
            raise CommandError(
                'Found workbooks but none matched known basin folder names. '
                f'Skipped {skipped} file(s).\n'
                'Folder names should include one of: '
                + '، '.join(spec.folder_key for spec in BASIN_SPECS)
            )

        self.stdout.write(self.style.SUCCESS(f'Done — {total_rows:,} observation rows processed.'))

    def _import_file(self, file_path: Path, basin: RainfallBasin) -> int:
        suffix = file_path.suffix.lower()
        if suffix == '.xls':
            return self._import_xls(file_path, basin)
        if suffix == '.xlsx':
            return self._import_xlsx(file_path, basin)
        return 0

    def _import_xls(self, file_path: Path, basin: RainfallBasin) -> int:
        if xlrd is None:
            raise CommandError(f'xlrd is required for .xls files: {_XLRD_IMPORT_ERROR}')
        workbook = xlrd.open_workbook(file_path)
        sheet = workbook.sheet_by_index(0)
        if sheet.nrows < 2:
            return 0

        def cell(row: int, col: int):
            if col >= sheet.ncols:
                return ''
            return sheet.cell_value(row, col)

        return self._import_rows(
            basin,
            row_count=sheet.nrows,
            cell=cell,
            xls_datemode=workbook.datemode,
        )

    def _import_xlsx(self, file_path: Path, basin: RainfallBasin) -> int:
        if openpyxl is None:
            raise CommandError(f'openpyxl is required for .xlsx files: {_OPENPYXL_IMPORT_ERROR}')
        workbook = openpyxl.load_workbook(file_path, read_only=True, data_only=True)
        try:
            sheet = workbook[workbook.sheetnames[0]]
            rows = list(sheet.iter_rows(values_only=True))
        finally:
            workbook.close()
        if len(rows) < 2:
            return 0

        def cell(row: int, col: int):
            values = rows[row]
            if col >= len(values):
                return ''
            return values[col]

        return self._import_rows(basin, row_count=len(rows), cell=cell, xls_datemode=None)

    def _import_rows(self, basin: RainfallBasin, *, row_count: int, cell, xls_datemode: int | None) -> int:
        batch: list[RainfallObservation] = []
        processed = 0

        for row_idx in range(1, row_count):
            station_name = str(cell(row_idx, 0) or '').strip()
            if not station_name:
                continue
            governorate = str(cell(row_idx, 1) or '').strip()
            obs_date = _parse_date(cell(row_idx, 2), xls_datemode=xls_datemode)
            if not obs_date:
                continue
            precipitation = _parse_decimal(cell(row_idx, 3))
            notes = str(cell(row_idx, 4) or '').strip()
            utm_x_raw = cell(row_idx, 5)
            utm_y_raw = cell(row_idx, 6)
            utm_x = _parse_decimal(utm_x_raw) if utm_x_raw not in (None, '') else None
            utm_y = _parse_decimal(utm_y_raw) if utm_y_raw not in (None, '') else None

            lat, lon = utm_to_wgs84(
                float(utm_x) if utm_x is not None else None,
                float(utm_y) if utm_y is not None else None,
            )

            station, _ = RainfallStation.objects.get_or_create(
                basin=basin,
                name=station_name,
                defaults={
                    'governorate': governorate,
                    'utm_x': utm_x,
                    'utm_y': utm_y,
                    'latitude': lat,
                    'longitude': lon,
                },
            )
            updated_fields: list[str] = []
            if governorate and station.governorate != governorate:
                station.governorate = governorate
                updated_fields.append('governorate')
            if lat is not None:
                if (
                    station.latitude != lat
                    or station.longitude != lon
                    or station.utm_x != utm_x
                    or station.utm_y != utm_y
                ):
                    station.latitude = lat
                    station.longitude = lon
                    station.utm_x = utm_x
                    station.utm_y = utm_y
                    updated_fields.extend(['latitude', 'longitude', 'utm_x', 'utm_y'])
            if updated_fields:
                station.save(update_fields=updated_fields)

            batch.append(
                RainfallObservation(
                    station=station,
                    observation_date=obs_date,
                    precipitation_mm=precipitation,
                    notes=notes,
                ),
            )
            processed += 1

            if len(batch) >= 2000:
                self._flush_batch(batch)
                batch.clear()

        if batch:
            self._flush_batch(batch)

        return processed

    @staticmethod
    def _flush_batch(batch: list[RainfallObservation]) -> None:
        with transaction.atomic():
            RainfallObservation.objects.bulk_create(
                batch,
                update_conflicts=True,
                unique_fields=['station', 'observation_date'],
                update_fields=['precipitation_mm', 'notes'],
            )
