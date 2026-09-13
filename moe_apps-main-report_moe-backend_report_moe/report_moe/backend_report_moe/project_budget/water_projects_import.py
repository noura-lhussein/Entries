"""Parse water-company projects workbook (targets, policies, quantitative targets)."""

from __future__ import annotations

from dataclasses import dataclass
from decimal import Decimal
from pathlib import Path

from openpyxl import load_workbook

from .budget_excel_import import clean_cell, normalize_name, parse_amount


@dataclass(frozen=True)
class ParsedWaterProject:
    row: int
    name_ar: str
    normalized_name: str
    company: str
    target: str
    policy: str
    approved_budget: Decimal | None
    quantitative_value: Decimal | None
    quantitative_unit: str | None


def parse_water_projects_workbook(
    path: Path,
    *,
    company_filter: str | None = None,
) -> list[ParsedWaterProject]:
    workbook = load_workbook(path, read_only=True, data_only=True)
    worksheet = workbook.active
    parsed: list[ParsedWaterProject] = []

    for row_index, raw_row in enumerate(worksheet.iter_rows(min_row=3, values_only=True), 3):
        cells = list(raw_row) + [None] * 8
        name = clean_cell(cells[3])
        if not name:
            continue

        company = clean_cell(cells[4])
        if company_filter and company != company_filter:
            continue

        qty_value = parse_amount(cells[7])
        parsed.append(
            ParsedWaterProject(
                row=row_index,
                name_ar=name,
                normalized_name=normalize_name(name),
                company=company,
                target=clean_cell(cells[1]),
                policy=clean_cell(cells[2]),
                approved_budget=parse_amount(cells[5]),
                quantitative_value=qty_value,
                quantitative_unit=clean_cell(cells[6]) or None,
            )
        )

    workbook.close()
    return parsed
