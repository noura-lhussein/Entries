"""Parse investment budget workbook into projects, milestones, and category codes."""

from __future__ import annotations

import re
from dataclasses import dataclass, field
from decimal import Decimal, InvalidOperation
from pathlib import Path

from openpyxl import load_workbook

BUDGET_SHEETS = (
    "المشاريع الجديدة",
    "المشاريع المباشر بها",
    "مشاريع الاستبدال والتجديد",
)

SECTION_CODES = frozenset({"31", "32", "33"})

SKIP_NAMES = frozenset(
    {
        "اسم المشروع",
        "الموازنة الاستثمارية لعام 2026",
        "النفقات الاستثمارية",
        "النفقات الاستثمارية ليرة جديدة",
        "النفقات الاستثمارية ليرة سورية",
        "المشاريع الجديدة",
        "المشاريع المباشر بها",
        "مشاريع الاستبدال والتجديد",
        "المجموع",
        "أراضي",
        "مباني وانشاءات ومرافق وطرق",
        "مباني وإنشاءات",
        "مباني وانشاءات",
        "آلات ومعدات",
        "وسائل نقل وانتقال",
        "عدد وأدوات وقوالب",
        "أثاث ومعدات مكاتب",
        "ثروة حيوانية ومائية",
        "نفقات تأسيس",
        "رواتب وأجور وتعويضات",
        "الآبار والسدود",
        "آبار مائية",
        "الآلات والمعدات",
        "آلات ومعدات مراكز الخدمات",
        "الكسوة",
    }
)


@dataclass(frozen=True)
class ParsedBudgetMilestone:
    sheet: str
    section: str
    row: int
    name_ar: str
    category_code: str | None
    category_name: str | None
    amount: Decimal | None
    project_group: str | None
    source_column: str  # "F" | "G"


@dataclass
class ParsedBudgetProjectGroup:
    name_ar: str
    section: str
    sheet: str
    milestones: list[ParsedBudgetMilestone] = field(default_factory=list)

    @property
    def total_amount(self) -> Decimal:
        total = Decimal("0")
        for item in self.milestones:
            if item.amount:
                total += item.amount
        return total


def normalize_name(value: str) -> str:
    text = " ".join(str(value).split()).strip()
    text = (
        text.replace("أ", "ا")
        .replace("إ", "ا")
        .replace("آ", "ا")
        .replace("ى", "ي")
        .replace("ة", "ه")
    )
    text = re.sub(r"[^\w\s/\-]", "", text)
    return text.lower()


def clean_cell(value) -> str:
    if value is None:
        return ""
    return " ".join(str(value).split()).strip()


def is_numeric_code(value) -> bool:
    if value is None:
        return False
    text = str(value).strip()
    if not text:
        return False
    return text.replace(".", "", 1).isdigit()


def parse_amount(value) -> Decimal | None:
    if value is None:
        return None
    if isinstance(value, (int, float)):
        return Decimal(str(value))
    text = str(value).strip().replace(",", "")
    if not text:
        return None
    try:
        return Decimal(text)
    except (InvalidOperation, ValueError):
        return None


def is_skipped_name(name: str) -> bool:
    if not name:
        return True
    if name in SKIP_NAMES:
        return True
    if len(name) <= 2:
        return True
    if is_numeric_code(name):
        return True
    return False


def is_project_header(name: str) -> bool:
    if is_skipped_name(name):
        return False
    if name.startswith("_"):
        return False
    if len(name) < 8:
        return False
    return True


def parse_budget_workbook(path: Path) -> tuple[list[ParsedBudgetMilestone], list[ParsedBudgetProjectGroup]]:
    workbook = load_workbook(path, read_only=True, data_only=True)
    milestones: list[ParsedBudgetMilestone] = []
    groups: dict[tuple[str, str], ParsedBudgetProjectGroup] = {}

    for sheet_name in BUDGET_SHEETS:
        if sheet_name not in workbook.sheetnames:
            continue
        worksheet = workbook[sheet_name]
        section = ""
        category_by_col: dict[int, tuple[str, str]] = {}
        project_group = ""

        for row_index, raw_row in enumerate(worksheet.iter_rows(min_row=4, values_only=True), 4):
            cells = list(raw_row) + [None] * 9

            if is_numeric_code(cells[1]) and str(cells[1]).strip() in SECTION_CODES:
                next_section = str(cells[1]).strip()
                if next_section != section:
                    category_by_col.clear()
                section = next_section
                project_group = ""

            for col in range(1, 5):
                if is_numeric_code(cells[col]):
                    code = str(cells[col]).strip().split(".")[0]
                    if code in SECTION_CODES:
                        continue
                    next_cell = clean_cell(cells[col + 1])
                    name = next_cell if next_cell and not is_numeric_code(
                        next_cell) else ""
                    for existing_col in list(category_by_col):
                        if existing_col > col:
                            del category_by_col[existing_col]
                    category_by_col[col] = (code, name)

            deepest_col = max(category_by_col) if category_by_col else None
            category_code, category_name = category_by_col.get(
                deepest_col, (None, ""))
            amount = parse_amount(cells[7])
            text_f = clean_cell(cells[5])
            text_g = clean_cell(cells[6])

            milestone_name = ""
            source_column = ""

            if text_g and not is_skipped_name(text_g):
                milestone_name = text_g
                source_column = "G"
            elif text_f and not is_skipped_name(text_f):
                if amount and amount > 0:
                    milestone_name = text_f
                    source_column = "F"
                elif category_code and len(category_code) >= 4:
                    milestone_name = text_f
                    source_column = "F"
                elif is_project_header(text_f):
                    project_group = text_f
                    key = (section or sheet_name, project_group)
                    groups.setdefault(
                        key,
                        ParsedBudgetProjectGroup(
                            name_ar=project_group,
                            section=section or "",
                            sheet=sheet_name,
                        ),
                    )
                elif len(text_f) >= 8:
                    milestone_name = text_f
                    source_column = "F"

            if not milestone_name:
                continue

            item = ParsedBudgetMilestone(
                sheet=sheet_name,
                section=section or "",
                row=row_index,
                name_ar=milestone_name,
                category_code=category_code,
                category_name=category_name or None,
                amount=amount if amount and amount > 0 else None,
                project_group=project_group or None,
                source_column=source_column,
            )
            milestones.append(item)
            if project_group:
                key = (section or sheet_name, project_group)
                group = groups.setdefault(
                    key,
                    ParsedBudgetProjectGroup(
                        name_ar=project_group,
                        section=section or "",
                        sheet=sheet_name,
                    ),
                )
                group.milestones.append(item)

    workbook.close()
    return milestones, list(groups.values())
