"""Parse budget category tree from docs/badjet/TREE.xlsx."""

from __future__ import annotations

from dataclasses import dataclass
from pathlib import Path

from openpyxl import load_workbook

# Excel columns (0-based): name, then codes from leaf-ish to root (جزئي → عام)
CODE_COLUMNS = (2, 3, 4, 5, 6)
NAME_COLUMN = 1
DESCRIPTION_COLUMN = 7


@dataclass(frozen=True)
class ParsedCategoryRow:
    code: str
    name_ar: str
    description: str


def _normalize_code(raw) -> str | None:
    if raw is None:
        return None
    text = str(raw).strip()
    if not text or not text.replace(".", "", 1).isdigit():
        return None
    return text.split(".")[0]


def _fix_known_code_typos(code: str) -> str:
    # Rows 42122–42126 belong under 2152 in the source tree (typo 4212 → 2152).
    if code.startswith("4212") and len(code) >= 5:
        return "2152" + code[4:]
    return code


def _next_sibling_code(code: str, used: set[str]) -> str:
    parent = code[:-1] if len(code) > 1 else ""
    length = len(code)
    siblings = [
        int(item)
        for item in used
        if len(item) == length and item.startswith(parent) and item.isdigit()
    ]
    next_value = (max(siblings) + 1) if siblings else int(code) + 1
    return str(next_value)


def parse_category_tree_xlsx(path: Path) -> list[ParsedCategoryRow]:
    workbook = load_workbook(path, read_only=True, data_only=True)
    worksheet = workbook.active
    used_codes: set[str] = set()
    parsed: list[ParsedCategoryRow] = []

    for row in worksheet.iter_rows(min_row=2, values_only=True):
        name_raw = row[NAME_COLUMN]
        if not name_raw:
            continue
        name_ar = str(name_raw).strip()
        if not name_ar or name_ar == "البيان":
            continue

        code = None
        for column in CODE_COLUMNS:
            code = _normalize_code(row[column])
            if code:
                break
        if not code:
            continue

        code = _fix_known_code_typos(code)
        while code in used_codes:
            code = _next_sibling_code(code, used_codes)

        description_raw = row[DESCRIPTION_COLUMN]
        description = str(description_raw).strip() if description_raw else ""

        used_codes.add(code)
        parsed.append(
            ParsedCategoryRow(code=code, name_ar=name_ar,
                              description=description)
        )

    workbook.close()
    return parsed


def build_category_tree_rows(
    parsed: list[ParsedCategoryRow],
) -> list[tuple[ParsedCategoryRow, str | None]]:
    by_code = {item.code: item for item in parsed}

    for item in list(by_code.values()):
        parent_code = item.code[:-1] if len(item.code) > 1 else None
        while parent_code and parent_code not in by_code:
            by_code[parent_code] = ParsedCategoryRow(
                code=parent_code,
                name_ar=f"فئة {parent_code}",
                description="",
            )
            parent_code = parent_code[:-1] if len(parent_code) > 1 else None

    return [
        (
            by_code[code],
            code[:-1] if len(code) > 1 else None,
        )
        for code in sorted(by_code.keys(), key=lambda value: (len(value), value))
    ]


# النفقات الاستثمارية (3) — skip these; use assistant codes (311, 312, …) as roots.
INVESTMENT_SECTION_SKIP = frozenset({"3", "31", "32", "33"})


def filter_investment_assistant_roots(
    tree_rows: list[tuple[ParsedCategoryRow, str | None]],
) -> list[tuple[ParsedCategoryRow, str | None]]:
    """Keep only codes under section 3; 311/312/313/… become top-level parents."""
    by_code = {item.code: (item, parent) for item, parent in tree_rows}

    selected = {
        code
        for code in by_code
        if code.startswith("3")
        and code not in INVESTMENT_SECTION_SKIP
        and len(code) >= 3
    }

    filtered: list[tuple[ParsedCategoryRow, str | None]] = []
    for code in sorted(selected, key=lambda value: (len(value), value)):
        item, parent_code = by_code[code]
        if len(code) == 3:
            filtered.append((item, None))
            continue
        while parent_code and parent_code not in selected:
            parent_code = by_code.get(parent_code, (None, None))[1]
        filtered.append((item, parent_code))
    return filtered
