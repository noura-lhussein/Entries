"""Extract official oil & gas daily Excel → structured dict for Info import."""

from __future__ import annotations

from datetime import date, datetime
from pathlib import Path
from typing import Any

from openpyxl import load_workbook

from .metric_catalog import (
    DAILY_FUEL_SALES_LABELS,
    EXECUTIVE_METRIC_SPECS,
    resolve_metric_labels,
)
from .report_template import (
    SHEET_FIELDS,
    SHEET_METRICS,
    SHEET_REFINERIES,
    SHEET_SUMMARY,
)


def _s(value: Any) -> str:
    if value is None:
        return ""
    if isinstance(value, datetime):
        return value.date().isoformat()
    if isinstance(value, date):
        return value.isoformat()
    return str(value).strip()


def _norm(value: Any) -> str:
    return " ".join(_s(value).lower().split())


def _parse_date(value: Any) -> str:
    if isinstance(value, datetime):
        return value.date().isoformat()
    if isinstance(value, date):
        return value.isoformat()
    text = _s(value)
    if not text:
        return ""
    for fmt in ("%Y-%m-%d", "%d/%m/%Y", "%d-%m-%Y", "%Y/%m/%d"):
        try:
            return datetime.strptime(text, fmt).date().isoformat()
        except ValueError:
            continue
    # Excel serial already coerced by data_only; leftover strings
    if len(text) >= 10 and text[4] == "-" and text[7] == "-":
        return text[:10]
    return ""


def _parse_number(value: Any) -> str:
    if value is None or value == "":
        return ""
    if isinstance(value, (int, float)):
        return str(value)
    text = _s(value).replace(",", "").replace("٫", ".")
    return text


def _metric_aliases(report_date: date) -> dict[str, str]:
    """Normalized label → metric_key (AR/EN, including Thu–Fri daily variants)."""
    aliases: dict[str, str] = {}
    for spec in EXECUTIVE_METRIC_SPECS:
        en, ar = resolve_metric_labels(spec, report_date)
        for label in (en, ar, spec.label_en, spec.label_ar):
            aliases[_norm(label)] = spec.key
        if spec.key in DAILY_FUEL_SALES_LABELS:
            den, dar = DAILY_FUEL_SALES_LABELS[spec.key]
            aliases[_norm(den)] = spec.key
            aliases[_norm(dar)] = spec.key
    return aliases


def is_official_oil_gas_workbook(path: Path) -> bool:
    wb = load_workbook(path, read_only=True, data_only=True)
    names = set(wb.sheetnames)
    wb.close()
    return SHEET_SUMMARY in names and SHEET_METRICS in names


def extract_xlsx(path: Path) -> dict[str, Any]:
    """Return report_date, national metrics, optional field/refinery entity rows."""
    if not is_official_oil_gas_workbook(path):
        raise ValueError(
            f'Workbook must include sheets "{SHEET_SUMMARY}" and "{SHEET_METRICS}".'
        )

    wb = load_workbook(path, data_only=True)
    try:
        summary_rows = [
            list(row) for row in wb[SHEET_SUMMARY].iter_rows(values_only=True)
        ]
        report_date = ""
        notes_ar = ""
        notes_en = ""
        for row in summary_rows:
            if not row:
                continue
            label = _norm(row[0])
            val = row[1] if len(row) > 1 else None
            if "report date" in label or "تاريخ التقرير" in label:
                report_date = _parse_date(val)
            elif label.startswith("ملاحظات") or "(ع)" in _s(row[0]).lower():
                notes_ar = _s(val)
            elif label in ("notes (en)", "notes"):
                notes_en = _s(val)

        if not report_date:
            raise ValueError("تعذّر قراءة تاريخ التقرير من ورقة الملخص.")

        try:
            d = date.fromisoformat(report_date)
        except ValueError as exc:
            raise ValueError(f"تاريخ تقرير غير صالح: {report_date}") from exc

        aliases = _metric_aliases(d)
        metrics: dict[str, str] = {"report_date": report_date}
        if notes_ar:
            metrics["notes_ar"] = notes_ar
        if notes_en:
            metrics["notes_en"] = notes_en

        metric_rows = list(wb[SHEET_METRICS].iter_rows(values_only=True))
        for row in metric_rows[1:]:
            if not row or all(c is None or _s(c) == "" for c in row):
                continue
            label_ar = _s(row[0]) if len(row) > 0 else ""
            label_en = _s(row[1]) if len(row) > 1 else ""
            value = _parse_number(row[2] if len(row) > 2 else None)
            if not value:
                continue
            key = aliases.get(_norm(label_ar)) or aliases.get(_norm(label_en))
            if key:
                metrics[key] = value

        fields: list[dict[str, Any]] = []
        if SHEET_FIELDS in wb.sheetnames:
            for row in list(wb[SHEET_FIELDS].iter_rows(values_only=True))[1:]:
                if not row or all(c is None or _s(c) == "" for c in row):
                    continue
                name_ar = _s(row[0]) if len(row) > 0 else ""
                name_en = _s(row[1]) if len(row) > 1 else ""
                if not name_ar and not name_en:
                    continue
                fields.append(
                    {
                        "name_ar": name_ar,
                        "name_en": name_en,
                        "crude_oil_bbl": _parse_number(row[2] if len(row) > 2 else None),
                        "natural_gas_mmscf": _parse_number(
                            row[3] if len(row) > 3 else None
                        ),
                        "condensate_bbl": _parse_number(row[4] if len(row) > 4 else None),
                    }
                )

        refineries: list[dict[str, Any]] = []
        if SHEET_REFINERIES in wb.sheetnames:
            for row in list(wb[SHEET_REFINERIES].iter_rows(values_only=True))[1:]:
                if not row or all(c is None or _s(c) == "" for c in row):
                    continue
                name_ar = _s(row[0]) if len(row) > 0 else ""
                name_en = _s(row[1]) if len(row) > 1 else ""
                if not name_ar and not name_en:
                    continue
                refineries.append(
                    {
                        "name_ar": name_ar,
                        "name_en": name_en,
                        "gasoline_ton": _parse_number(row[2] if len(row) > 2 else None),
                        "diesel_ton": _parse_number(row[3] if len(row) > 3 else None),
                        "fuel_oil_ton": _parse_number(row[4] if len(row) > 4 else None),
                        "lpg_ton": _parse_number(row[5] if len(row) > 5 else None),
                    }
                )

        return {
            "report_date": report_date,
            "metrics": metrics,
            "fields": fields,
            "refineries": refineries,
        }
    finally:
        wb.close()
