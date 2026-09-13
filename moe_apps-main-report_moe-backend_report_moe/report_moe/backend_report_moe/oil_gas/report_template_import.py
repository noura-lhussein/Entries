"""Blank Excel template and import for oil & gas executive daily reports."""

from __future__ import annotations

import tempfile
from datetime import date, datetime
from decimal import Decimal, InvalidOperation
from pathlib import Path
from typing import Any

from django.http import HttpResponse
from openpyxl import load_workbook
from water.template_generator import multisheet_xlsx_response

from .metric_catalog import EXECUTIVE_METRIC_SPECS, METRIC_BY_KEY, resolve_metric_labels
from .models import DailyReport
from .services import upsert_metric

SHEET_INFO = 'معلومات'
SHEET_METRICS = 'المؤشرات'

INFO_DATE_LABELS = ('تاريخ التقرير', 'report_date', 'Report date')
INFO_STATUS_LABELS = ('الحالة', 'status', 'Status')
INFO_NOTES_AR_LABELS = ('ملاحظات', 'notes_ar', 'Notes')


def build_oil_gas_executive_template_sheets(
    sample_date: date | None = None,
) -> dict[str, list[list[Any]]]:
    report_date = sample_date or date.today()
    info = [
        ['الحقل', 'القيمة'],
        ['تاريخ التقرير', report_date.isoformat()],
        ['الحالة', 'published'],
        ['ملاحظات', ''],
    ]
    metrics: list[list[Any]] = [['المفتاح', 'البيان', 'الوحدة', 'القيمة']]
    for spec in EXECUTIVE_METRIC_SPECS:
        _en, label_ar = resolve_metric_labels(spec, report_date)
        metrics.append([spec.key, label_ar, spec.unit, None])
    return {SHEET_INFO: info, SHEET_METRICS: metrics}


def export_oil_gas_executive_template_xlsx(sample_date: date | None = None) -> HttpResponse:
    sheets = build_oil_gas_executive_template_sheets(sample_date)
    return multisheet_xlsx_response('oil_gas_executive_report_template.xlsx', sheets)


def _cell_str(value: Any) -> str:
    if value is None:
        return ''
    if isinstance(value, datetime):
        return value.date().isoformat()
    if isinstance(value, date):
        return value.isoformat()
    return str(value).strip()


def _parse_date(value: Any) -> date | None:
    if isinstance(value, datetime):
        return value.date()
    if isinstance(value, date):
        return value
    text = _cell_str(value)
    if not text:
        return None
    for fmt in ('%Y-%m-%d', '%d/%m/%Y', '%d-%m-%Y', '%Y/%m/%d'):
        try:
            return datetime.strptime(text, fmt).date()
        except ValueError:
            continue
    return None


def _parse_decimal(value: Any) -> Decimal | None:
    if value is None or value == '':
        return None
    if isinstance(value, (int, float, Decimal)):
        return Decimal(str(value))
    text = str(value).strip().replace(',', '').replace('٫', '.')
    if not text:
        return None
    try:
        return Decimal(text)
    except (InvalidOperation, ValueError):
        return None


def _info_map(rows: list[list[Any]]) -> dict[str, Any]:
    mapping: dict[str, Any] = {}
    for row in rows[1:]:
        if not row:
            continue
        key = _cell_str(row[0]).lower()
        val = row[1] if len(row) > 1 else None
        mapping[key] = val
    return mapping


def _lookup_info(info: dict[str, Any], labels: tuple[str, ...]) -> Any:
    for label in labels:
        for key, value in info.items():
            if key == label.lower() or label.lower() in key:
                return value
    return None


def parse_oil_gas_executive_xlsx(path: Path) -> dict[str, Any]:
    workbook = load_workbook(path, data_only=True)
    if SHEET_INFO not in workbook.sheetnames or SHEET_METRICS not in workbook.sheetnames:
        raise ValueError(
            f'Template must include sheets "{SHEET_INFO}" and "{SHEET_METRICS}".'
        )

    info_rows = [list(row) for row in workbook[SHEET_INFO].iter_rows(values_only=True)]
    metric_rows = [list(row) for row in workbook[SHEET_METRICS].iter_rows(values_only=True)]
    info = _info_map(info_rows)

    report_date = _parse_date(_lookup_info(info, INFO_DATE_LABELS))
    if report_date is None:
        raise ValueError('Report date is required (YYYY-MM-DD) in sheet معلومات.')

    status_raw = _cell_str(_lookup_info(info, INFO_STATUS_LABELS)).lower()
    if status_raw in ('draft', 'مسودة'):
        status = DailyReport.Status.DRAFT
    else:
        status = DailyReport.Status.PUBLISHED

    notes_ar = _cell_str(_lookup_info(info, INFO_NOTES_AR_LABELS))

    label_to_key = {
        resolve_metric_labels(spec, report_date)[1].strip(): spec.key
        for spec in EXECUTIVE_METRIC_SPECS
    }

    metrics: list[dict[str, Any]] = []
    for row in metric_rows[1:]:
        if not row or all(cell is None or _cell_str(cell) == '' for cell in row):
            continue
        key = _cell_str(row[0])
        label = _cell_str(row[1]) if len(row) > 1 else ''
        unit = _cell_str(row[2]) if len(row) > 2 else ''
        value = _parse_decimal(row[3] if len(row) > 3 else None)
        if value is None:
            continue
        if key not in METRIC_BY_KEY:
            key = label_to_key.get(label, '')
        if not key or key not in METRIC_BY_KEY:
            continue
        spec = METRIC_BY_KEY[key]
        metrics.append(
            {
                'metric_key': key,
                'dimension': '',
                'value': value,
                'unit': unit or spec.unit,
            }
        )

    if not metrics:
        raise ValueError('No metric values found in sheet المؤشرات.')

    return {
        'report_date': report_date,
        'status': status,
        'notes_ar': notes_ar,
        'notes_en': '',
        'metrics': metrics,
    }


def import_oil_gas_executive_xlsx(
    path: Path,
    *,
    user=None,
    publish: bool | None = None,
) -> DailyReport:
    payload = parse_oil_gas_executive_xlsx(path)
    if publish is True:
        payload['status'] = DailyReport.Status.PUBLISHED
    elif publish is False:
        payload['status'] = DailyReport.Status.DRAFT

    report, _ = DailyReport.objects.update_or_create(
        report_date=payload['report_date'],
        defaults={
            'status': payload['status'],
            'notes_ar': payload.get('notes_ar', ''),
            'notes_en': payload.get('notes_en', ''),
            'created_by': user,
        },
    )
    for item in payload['metrics']:
        upsert_metric(
            report,
            item['metric_key'],
            item['value'],
            item.get('dimension') or '',
            item.get('unit') or '',
        )
    return report


def import_oil_gas_executive_upload(
    uploaded,
    *,
    user=None,
    publish: bool = True,
) -> dict[str, Any]:
    suffix = Path(getattr(uploaded, 'name', '') or 'upload.xlsx').suffix.lower() or '.xlsx'
    if suffix not in ('.xlsx', '.xlsm'):
        return {'ok': False, 'error': 'Only .xlsx files are supported.'}
    with tempfile.NamedTemporaryFile(suffix=suffix, delete=False) as tmp:
        for chunk in uploaded.chunks():
            tmp.write(chunk)
        tmp_path = Path(tmp.name)
    try:
        report = import_oil_gas_executive_xlsx(tmp_path, user=user, publish=publish)
    except Exception as exc:  # noqa: BLE001 — surface parse errors to API
        return {'ok': False, 'error': str(exc)}
    finally:
        tmp_path.unlink(missing_ok=True)

    return {
        'ok': True,
        'id': report.id,
        'report_date': report.report_date.isoformat(),
        'status': report.status,
        'metrics_count': report.metrics.count(),
    }
